# -*- coding: utf-8 -*-
"""OpenRotas Desktop — testes da edição desktop (Etapa 6).

Cobre a lógica que NÃO depende de Windows/GUI/OSRM: configuração, perfil de desempenho,
resolução do motor (fallback defensivo), camada de dados locais e diagnóstico. Roda em
qualquer ambiente (inclusive este, Linux na nuvem) com `pytest desktop/tests`."""
import os
import sys
import pytest

_DESKTOP = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(_DESKTOP, "app"))
sys.path.insert(0, _DESKTOP)
sys.path.insert(0, os.path.join(_DESKTOP, "data_local"))
sys.path.insert(0, os.path.join(_DESKTOP, "telemetry"))
sys.path.insert(0, os.path.join(_DESKTOP, "installer"))

import desktop_config as cfg          # noqa: E402
import diagnostics                    # noqa: E402
import janela                         # noqa: E402
from engines import osrm_manager as osrm   # noqa: E402
import local_data                     # noqa: E402
from resources import resource_manager as rm   # noqa: E402
from resources import painel                    # noqa: E402
from resources import provisionamento as prov   # noqa: E402
from resources import central_dados             # noqa: E402
from audit import coverage_auditor as ca         # noqa: E402
from geo import repositorio as georepo            # noqa: E402
import exec_profile                   # noqa: E402
import app_update                     # noqa: E402


# ------------------------------- desktop_config -------------------------------
def test_app_root_tem_o_streamlit_app():
    assert (cfg.app_root() / "streamlit_app.py").exists()

def test_hardware_coerente():
    hw = cfg.detectar_hardware()
    assert hw["cpus"] >= 1
    assert hw["perfil"] in ("performance", "equilibrado", "economico")

def test_perfil_desempenho_dentro_dos_limites():
    p = cfg.perfil_desempenho()
    assert 8 <= p["workers_rota"] <= 32
    assert 512 <= p["cache_mb_sugerido"] <= 4096

def test_validar_config_service_role_e_erro():
    av = cfg.validar_config({"local_login": True, "SUPABASE_ANON_KEY": "eyJ...service_role...x"})
    assert any(n == "erro" and "service_role" in m for n, m in av)

def test_validar_config_docker_sem_grafo_avisa():
    av = cfg.validar_config({"local_login": True, "osrm": {"mode": "docker"}})
    assert any(n == "aviso" and "graph_url" in m for n, m in av)

def test_validar_config_supabase_ausente_sem_login_local():
    av = cfg.validar_config({"osrm": {"mode": "off"}})
    assert any(n == "aviso" and "SUPABASE" in m for n, m in av)

def test_validar_config_login_local_sem_supabase_ok():
    av = cfg.validar_config({"local_login": True, "osrm": {"mode": "off"}})
    assert not any("SUPABASE" in m for _, m in av)   # login local não exige Supabase

def test_validar_config_graph_url_com_mode_off_e_info():
    av = cfg.validar_config({"local_login": True, "osrm": {"mode": "off", "graph_url": "http://x/g.part00"}})
    assert any(n == "info" and "graph_url" in m for n, m in av)


def test_user_data_dir_e_resumo():
    paths = cfg.ensure_user_dirs()
    assert paths["cache"].exists()
    r = cfg.resumo_config()
    assert {"app_root", "hardware", "porta", "user_data"} <= set(r)

def test_bootstrap_config_cria_do_exemplo(tmp_path, monkeypatch):
    # 1º uso: cria desktop.json a partir do exemplo; 2ª chamada é no-op.
    monkeypatch.setattr(cfg, "user_data_dir", lambda: tmp_path)
    criado = cfg.bootstrap_config_usuario()
    destino = tmp_path / "config" / "desktop.json"
    assert criado == destino and destino.exists()
    import json as _json
    dados = _json.loads(destino.read_text(encoding="utf-8"))
    assert "SUPABASE_URL" in dados and "osrm" in dados     # veio do exemplo
    assert cfg.bootstrap_config_usuario() is None          # já existe → não recria

def test_caminho_exemplo_config_existe():
    assert cfg._caminho_exemplo_config() is not None       # o exemplo embarcado é encontrável


# ------------------------------- osrm_manager --------------------------------
def test_resolver_off_nao_mexe_na_url():
    r = osrm.resolver({"mode": "off"})
    assert r.modo == "off" and r.url is None and r.ativo is False

def test_resolver_external_sem_servidor_nao_quebra():
    r = osrm.resolver({"mode": "external", "url": "http://127.0.0.1:5999"})
    assert r.modo == "external" and r.ativo is False and r.url is None  # fallback ao público

def test_resolver_legado_osrm_url_tratado_como_external():
    r = osrm.resolver(None, osrm_url_legado="http://127.0.0.1:5999")
    assert r.modo == "external"

def test_resolver_modo_desconhecido_defensivo():
    r = osrm.resolver({"mode": "xpto"})
    assert r.ativo is False and r.url is None

def test_health_url_vazia_e_encerrar_noop():
    assert osrm.health("") is False
    osrm.encerrar(osrm.ResultadoMotor(url=None, modo="off", ativo=False))  # não levanta

def test_remover_container_nunca_levanta():
    osrm._remover_container("")                    # nome vazio → no-op
    osrm._remover_container("openrotas-osrm-inexistente")   # docker ausente → não levanta

def test_encerrar_gerenciado_sem_docker_nao_levanta():
    # ResultadoMotor gerenciado com nome: encerrar tenta parar o container (docker ausente) sem erro
    r = osrm.ResultadoMotor(url="http://localhost:5000", modo="docker", ativo=True,
                            gerenciado=True, nome=osrm.CONTAINER_OSRM)
    osrm.encerrar(r)


def test_garantir_grafo_sem_url_e_noop(tmp_path):
    # sem graph_path e sem graph_url → nada a provisionar, devolve None (cai no público)
    assert osrm.garantir_grafo({}, tmp_path) is None

def test_garantir_grafo_usa_path_existente(tmp_path):
    g = tmp_path / "brazil-latest.osrm"
    g.write_text("x")
    assert osrm.garantir_grafo({"graph_path": str(g)}, tmp_path / "outro") == str(g)

def test_garantir_grafo_reusa_ja_provisionado(tmp_path):
    (tmp_path / "brazil-latest.osrm").write_text("x")  # já provisionado antes
    got = osrm.garantir_grafo({"graph_url": "http://exemplo/inexistente.tar.gz"}, tmp_path)
    assert got == str(tmp_path / "brazil-latest.osrm")  # não baixa; reusa

def test_base_partes_normaliza():
    assert osrm._base_partes("http://x/g.tar.gz.part") == "http://x/g.tar.gz.part"
    assert osrm._base_partes("http://x/g.tar.gz.part00") == "http://x/g.tar.gz.part"
    assert osrm._base_partes("http://x/g.tar.gz.part7") == "http://x/g.tar.gz.part"

def test_garantir_grafo_em_partes(tmp_path):
    import io, tarfile, gzip
    # cria um tar.gz com um 'brazil-latest.osrm' dentro
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w") as t:
        dado = b"GRAFO-OSRM-FAKE" * 5000
        info = tarfile.TarInfo("brazil-latest.osrm"); info.size = len(dado)
        t.addfile(info, io.BytesIO(dado))
    targz = gzip.compress(buf.getvalue())
    # divide em 2 partes
    meio = len(targz) // 2
    srv = tmp_path / "srv"; srv.mkdir()
    (srv / "brazil-osrm-mld.tar.gz.part00").write_bytes(targz[:meio])
    (srv / "brazil-osrm-mld.tar.gz.part01").write_bytes(targz[meio:])
    base_uri = (srv / "brazil-osrm-mld.tar.gz.part").as_uri()
    destino = tmp_path / "dl"
    got = osrm.garantir_grafo({"graph_url": base_uri}, destino)
    assert got == str(destino / "brazil-latest.osrm") and (destino / "brazil-latest.osrm").exists()

def test_baixar_partes_sem_nenhuma_parte(tmp_path):
    base = (tmp_path / "naoexiste.part").as_uri()
    assert osrm._baixar_partes(base, tmp_path / "out.tar.gz") is False
    assert not (tmp_path / "out.tar.gz").exists()

def test_baixar_partes_rejeita_parte_com_tamanho_errado(tmp_path, monkeypatch):
    # Correção do "pacote sem .osrm": uma parte corrompida (tamanho != publicado) deve ser
    # DETECTADA, apagada e o download abortado — em vez de concatenar lixo.
    srv = tmp_path / "srv"; srv.mkdir()
    (srv / "brazil-osrm-mld.tar.gz.part00").write_bytes(b"A" * 1000)
    (srv / "brazil-osrm-mld.tar.gz.part01").write_bytes(b"B" * 500)   # "corrompida": tamanho menor
    base_uri = (srv / "brazil-osrm-mld.tar.gz.part").as_uri()
    # tamanhos "publicados": part00 bate (1000); part01 deveria ter 777 (não bate com 500)
    monkeypatch.setattr(osrm, "_tamanhos_esperados",
                        lambda _u: {"brazil-osrm-mld.tar.gz.part00": 1000,
                                    "brazil-osrm-mld.tar.gz.part01": 777})
    out = tmp_path / "out.tar.gz"
    ok = osrm._baixar_partes(base_uri, out, progresso=lambda ev: None)
    assert ok is False                                   # abortou por integridade
    assert not out.exists()                              # não concatenou lixo

def test_instalar_grafo_de_partes_locais(tmp_path):
    # Caminho GARANTIDO (sem rede): usuário baixou as partes pelo navegador numa pasta.
    import io, tarfile, gzip
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w") as t:
        dado = b"GRAFO-OSRM-FAKE" * 5000
        info = tarfile.TarInfo("brazil-latest.osrm"); info.size = len(dado)
        t.addfile(info, io.BytesIO(dado))
    targz = gzip.compress(buf.getvalue())
    meio = len(targz) // 2
    pasta = tmp_path / "baixados"; pasta.mkdir()
    (pasta / "brazil-osrm-mld.tar.gz.part00").write_bytes(targz[:meio])
    (pasta / "brazil-osrm-mld.tar.gz.part01").write_bytes(targz[meio:])
    destino = tmp_path / "dados"
    got = osrm.instalar_grafo_de_partes_locais(pasta, destino)
    assert got == str(destino / "brazil-latest.osrm") and (destino / "brazil-latest.osrm").exists()

def test_instalar_grafo_de_partes_locais_pasta_vazia(tmp_path):
    vazia = tmp_path / "vazia"; vazia.mkdir()
    assert osrm.instalar_grafo_de_partes_locais(vazia, tmp_path / "dados") is None

def test_garantir_grafo_nome_base_diferente(tmp_path):
    # o move deve derivar o padrão do arquivo encontrado, não de 'brazil-latest' fixo.
    import io, tarfile, gzip
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w") as t:
        for nome in ("brasil.osrm", "brasil.osrm.edges"):
            dado = (nome + "-data").encode() * 100
            info = tarfile.TarInfo(nome); info.size = len(dado)
            t.addfile(info, io.BytesIO(dado))
    (tmp_path / "g.tar.gz").write_bytes(gzip.compress(buf.getvalue()))
    destino = tmp_path / "dl"
    got = osrm.garantir_grafo({"graph_url": (tmp_path / "g.tar.gz").as_uri()}, destino)
    assert got == str(destino / "brasil.osrm")
    assert (destino / "brasil.osrm.edges").exists()       # artefato irmão também foi movido

def test_aquecer_motor_best_effort_nao_levanta():
    # Aquecimento é oportunista: URL vazia → 0; URL que recusa conexão → 0, sem levantar.
    assert osrm._aquecer_motor("", timeout=1.0) == 0
    assert osrm._aquecer_motor("http://127.0.0.1:1", rotas=("-46.6,-23.5;-43.1,-22.9",),
                               timeout=1.0) == 0

def test_aquecer_motor_conta_rotas_ok(monkeypatch):
    # Com health/rota respondendo 'Ok', conta cada rota aquecida.
    class _Resp:
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def read(self, *a): return b'{"code":"Ok"}'
    monkeypatch.setattr(osrm.urllib.request, "urlopen", lambda *a, **k: _Resp())
    n = osrm._aquecer_motor("http://127.0.0.1:5000",
                            rotas=("-46.6,-23.5;-43.1,-22.9", "-47.9,-15.8;-49.2,-16.7"), timeout=1.0)
    assert n == 2

def test_autoligar_se_tem_grafo(tmp_path):
    # TURNKEY: modo desligado + grafo presente → 'auto'; sem grafo → mantém o modo.
    vazia = tmp_path / "vazia"; vazia.mkdir()
    com = tmp_path / "com"; com.mkdir()
    (com / "brazil-latest.osrm.mldgr").write_bytes(b"x")     # família presente
    assert osrm.autoligar_se_tem_grafo("off", vazia) == "off"
    assert osrm.autoligar_se_tem_grafo("", vazia) == "off"
    assert osrm.autoligar_se_tem_grafo("off", vazia, com) == "auto"   # acha em alguma pasta
    assert osrm.autoligar_se_tem_grafo("", com) == "auto"
    # respeita um modo já escolhido pelo usuário (não rebaixa nem sobrescreve)
    assert osrm.autoligar_se_tem_grafo("local", vazia) == "local"
    assert osrm.autoligar_se_tem_grafo("docker", com) == "docker"

def test_base_osrm_em_familia_sem_arquivo_base(tmp_path):
    # CAUSA RAIZ: OSRM 5.x (MLD) não gera um '<base>.osrm' literal — só a família '<base>.osrm.*'.
    # _base_osrm_em deve derivar o caminho-base a partir da família.
    for suf in (".cells", ".mldgr"):
        (tmp_path / ("brazil-latest.osrm" + suf)).write_bytes(b"x")
    assert osrm._base_osrm_em(tmp_path) == tmp_path / "brazil-latest.osrm"
    # com um arquivo-base AVULSO presente → devolve-o (compat com formatos antigos)
    outro = tmp_path / "sub"; outro.mkdir()
    (outro / "g.osrm").write_bytes(b"x")
    assert osrm._base_osrm_em(outro) == outro / "g.osrm"
    # pasta vazia / inexistente → None (não levanta)
    vaz = tmp_path / "v"; vaz.mkdir()
    assert osrm._base_osrm_em(vaz) is None
    assert osrm._base_osrm_em(tmp_path / "nao_existe") is None

def test_garantir_grafo_familia_mld_sem_arquivo_base(tmp_path):
    # Caso REAL do grafo do Brasil: o pacote NÃO contém 'brazil-latest.osrm', só a família
    # 'brazil-latest.osrm.*'. Antes da correção isso dava "pacote do grafo sem arquivo .osrm"
    # em TODO download, mesmo íntegro. Agora deve provisionar e devolver o caminho-BASE.
    import io, tarfile, gzip
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w") as t:
        for suf in (".cells", ".mldgr", ".partition", ".properties"):
            nome = "brazil-latest.osrm" + suf
            dado = (nome + "-data").encode() * 50
            info = tarfile.TarInfo(nome); info.size = len(dado)
            t.addfile(info, io.BytesIO(dado))
    (tmp_path / "g.tar.gz").write_bytes(gzip.compress(buf.getvalue()))
    destino = tmp_path / "dl"
    got = osrm.garantir_grafo({"graph_url": (tmp_path / "g.tar.gz").as_uri()}, destino)
    assert got == str(destino / "brazil-latest.osrm")        # caminho-base (sem sufixo)
    assert not (destino / "brazil-latest.osrm").exists()     # não existe como arquivo literal
    assert (destino / "brazil-latest.osrm.mldgr").exists()   # a família inteira foi movida
    assert (destino / "brazil-latest.osrm.cells").exists()
    assert (destino / "brazil-latest.osrm.partition").exists()
    # 2ª chamada reusa o já provisionado (reconhece a família, não rebaixa)
    assert osrm.garantir_grafo({"graph_url": "http://x/inexistente.tar.gz"}, destino) == got

def test_instalar_grafo_familia_mld_sem_arquivo_base(tmp_path):
    # Mesmo caso REAL pelo caminho OFFLINE ("Importar mapa de uma pasta…").
    import io, tarfile, gzip
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w") as t:
        for suf in (".cells", ".mldgr", ".partition"):
            nome = "brazil-latest.osrm" + suf
            dado = (nome + "-data").encode() * 50
            info = tarfile.TarInfo(nome); info.size = len(dado)
            t.addfile(info, io.BytesIO(dado))
    targz = gzip.compress(buf.getvalue())
    meio = len(targz) // 2
    pasta = tmp_path / "baixados"; pasta.mkdir()
    (pasta / "brazil-osrm-mld.tar.gz.part00").write_bytes(targz[:meio])
    (pasta / "brazil-osrm-mld.tar.gz.part01").write_bytes(targz[meio:])
    destino = tmp_path / "dados"
    got = osrm.instalar_grafo_de_partes_locais(pasta, destino)
    assert got == str(destino / "brazil-latest.osrm")
    assert (destino / "brazil-latest.osrm.mldgr").exists()

def test_subir_local_aceita_grafo_familia_sem_arquivo_base(tmp_path, monkeypatch):
    # O motor nativo deve ACEITAR o grafo pelo caminho-base quando só a família '<base>.osrm.*'
    # existe (não rejeitar como 'grafo não encontrado'). Sem binário → cai no público, mas a
    # mensagem prova que passou da checagem do grafo.
    for suf in (".cells", ".mldgr", ".partition"):
        (tmp_path / ("brazil-latest.osrm" + suf)).write_bytes(b"x")
    base = tmp_path / "brazil-latest.osrm"                    # NÃO existe como arquivo
    monkeypatch.setattr(osrm, "_bin_osrm_routed", lambda: None)
    r = osrm._subir_local({"mode": "local", "graph_path": str(base)})
    assert r.ativo is False
    assert "grafo não encontrado" not in (r.detalhe or "")    # passou da checagem de grafo
    assert "binário" in (r.detalhe or "")                     # parou na falta do binário

def test_remover_container_retorna_bool():
    assert osrm._remover_container("") is False            # nome vazio
    assert isinstance(osrm._remover_container("openrotas-osrm-x"), bool)


# -------------------------------- local_data ---------------------------------
@pytest.fixture(scope="module")
def reg():
    return local_data.LocalDataRegistry(cfg.app_root(), cfg.user_data_dir() / "data_local")

def test_essenciais_presentes(reg):
    assert reg.essenciais_ok() is True

def test_municipios_carrega_e_tem_linhas(reg):
    df = reg.municipios()
    assert len(df) > 5000  # ~5.570 municípios

def test_indice_ibge_nao_vazio(reg):
    idx = reg.indice_municipios_por_ibge()
    assert len(idx) > 5000

def test_assinatura_detecta_arquivo(reg):
    a = reg.assinatura("municipios")
    assert a["existe"] and a["bytes"] > 0 and len(a["hash12"]) == 12

def test_offline_pronto_estrutura(reg):
    off = reg.offline_pronto()
    assert off["geocodificacao_hidro_local"] is True      # bases embarcadas sempre presentes
    # roteamento_local/pronto refletem a presença REAL do grafo (alguns ambientes já o trazem).
    # Testar a invariante (== existe(osrm_brasil)) torna o teste portável sem afrouxar a honestidade:
    # quando o grafo falta, 'faltam' tem de dizê-lo.
    tem = reg.existe("osrm_brasil")
    assert off["roteamento_local"] is bool(tem)
    assert off["pronto"] is bool(tem)
    if not tem:
        assert "osrm_brasil" in " ".join(off["faltam"])

def test_osrm_brasil_e_opcional(reg):
    # osrm_brasil (grafo) é um recurso OPCIONAL: pode ou não estar no ambiente, mas nunca é
    # essencial (os essenciais são as bases embarcadas, sempre presentes). Invariante de
    # classificação — portável, não assume ausência num ambiente que já traga o grafo.
    g = next(d for d in local_data.CATALOGO if d.chave == "osrm_brasil")
    assert g.essencial is False
    assert reg.essenciais_ok() is True

def test_local_data_assinatura_familia_osrm(tmp_path):
    # Correção: assinatura() deve medir a FAMÍLIA '<base>.osrm.*' (o MLD não tem arquivo-base),
    # senão verificar()/status() marcam um grafo íntegro como corrompido/0 MB.
    r = local_data.LocalDataRegistry(cfg.app_root(), tmp_path / "data_local")
    dl = tmp_path / "data_local"; dl.mkdir(parents=True, exist_ok=True)
    (dl / "brazil-latest.osrm.cells").write_bytes(b"A" * 2048)
    (dl / "brazil-latest.osrm.mldgr").write_bytes(b"B" * 4096)
    a = r.assinatura("osrm_brasil")
    assert a["existe"] is True and a["bytes"] == 2048 + 4096 and a["hash12"]
    # base genuinamente ausente (sem família) → existe False
    r2 = local_data.LocalDataRegistry(cfg.app_root(), tmp_path / "vazio")
    assert r2.assinatura("osrm_brasil")["existe"] is False

def test_osrm_brasil_reconhecido_pela_familia_sem_arquivo_base(tmp_path):
    # OSRM 5.x (MLD) não cria 'brazil-latest.osrm' literal — só a família '<base>.osrm.*'.
    # existe()/legivel()/offline_pronto devem reconhecer o grafo pela FAMÍLIA (senão a Central
    # de Dados reportaria "sem grafo" num grafo recém-instalado e íntegro).
    r = local_data.LocalDataRegistry(cfg.app_root(), tmp_path / "data_local")
    dl = tmp_path / "data_local"; dl.mkdir(parents=True, exist_ok=True)
    assert r.existe("osrm_brasil") is False                   # nada ainda
    for suf in (".cells", ".mldgr", ".partition"):
        (dl / ("brazil-latest.osrm" + suf)).write_bytes(b"x")
    assert not (dl / "brazil-latest.osrm").exists()           # sem arquivo-base literal
    assert r.existe("osrm_brasil") is True                    # mas a família conta como instalado
    assert r.legivel("osrm_brasil") is True
    assert r.offline_pronto()["roteamento_local"] is True

def test_legivel_base_valida(reg):
    assert reg.legivel("municipios") is True          # parquet real abre (lê só o schema)

def test_legivel_detecta_parquet_corrompido(tmp_path):
    # um "parquet" com lixo deve ser detectado como não-legível (§18)
    r = local_data.LocalDataRegistry(cfg.app_root(), tmp_path / "data_local")
    ov = r.override_dir / "municipios.parquet"
    ov.parent.mkdir(parents=True, exist_ok=True)
    ov.write_bytes(b"isto nao e um parquet valido")
    assert r.existe("municipios") is True             # existe (override)
    assert r.legivel("municipios") is False           # mas não abre → corrompido

def test_override_vence_o_bundle(tmp_path):
    # Uma cópia reparada/atualizada no diretório de override do perfil do usuário tem
    # prioridade sobre a base embarcada read-only (§18/§19/§45).
    r = local_data.LocalDataRegistry(cfg.app_root(), tmp_path / "data_local")
    bundle = r.caminho("municipios")
    assert "data_local" not in str(bundle)          # sem override → aponta para o bundle
    ov = r.override_dir / "municipios.parquet"
    ov.parent.mkdir(parents=True, exist_ok=True)
    ov.write_text("reparado")
    assert r.caminho("municipios") == ov            # com override → vence o bundle


# ----------------------------- resource_manager ------------------------------
def test_rm_status_lista_recursos():
    linhas = rm.status()
    chaves = {r["chave"] for r in linhas}
    assert {"municipios", "osrm_brasil"} <= chaves
    mun = next(r for r in linhas if r["chave"] == "municipios")
    assert mun["instalado"] is True and mun["estado"] == rm.OK and mun["obrigatorio"] is True
    assert mun["modulo"]  # cada recurso sabe qual módulo o usa

def test_rm_osrm_opcional():
    g = next(r for r in rm.status() if r["chave"] == "osrm_brasil")
    # Classificação (invariante): provisionável e NÃO obrigatório.
    assert g["tipo"] == "provisionável" and g["obrigatorio"] is False
    # Estado reflete o ambiente real: instalado → estado diferente de OPCIONAL_AUSENTE;
    # ausente → OPCIONAL_AUSENTE. Portável (alguns ambientes já trazem o grafo).
    if g["instalado"]:
        assert g["estado"] != rm.OPCIONAL_AUSENTE
    else:
        assert g["estado"] == rm.OPCIONAL_AUSENTE

def test_rm_verificar_essenciais_ok():
    v = rm.verificar()
    assert v["ok"] is True and v["faltam_obrigatorios"] == []

def test_rm_provisionar_sem_url_e_noop():
    r = rm.provisionar_grafo({})
    # Sem graph_url/graph_path não há o que baixar. Se o grafo JÁ está instalado, garantir_grafo
    # devolve o caminho existente (no-op bem-sucedido); se não está, ok=False. Portável.
    g = next(x for x in rm.status() if x["chave"] == "osrm_brasil")
    assert r["ok"] is bool(g["instalado"])

def test_rm_resumo_ambiente_texto():
    txt = rm.resumo_ambiente()
    assert "Recursos do software" in txt and "municipios" in txt


# ---------------------- painel (Central de Recursos — §17/§24) ---------------
def test_painel_html_tem_secoes_e_recursos():
    h = painel.construir_html()
    assert "<!doctype html>" in h and "Central de Recursos" in h
    assert "Recursos do software" in h and "Prontidão offline" in h and "Perfil de execução" in h
    assert "municipios" in h          # a tabela de recursos foi preenchida
    assert "prefers-color-scheme" in h  # tema claro/escuro

def test_painel_gera_arquivo(tmp_path):
    out = tmp_path / "central.html"
    got = painel.gerar(str(out))
    assert got == str(out) and out.exists()
    conteudo = out.read_text(encoding="utf-8")
    assert conteudo.startswith("<!doctype html>") and "</html>" in conteudo

def test_painel_escapa_conteudo_sem_quebrar():
    # construir_html nunca levanta e produz HTML bem-formado mesmo sem osrm_cfg
    h = painel.construir_html(osrm_cfg={"graph_url": "http://x/<b>"})
    assert "<html" in h and "</html>" in h


# ------------------------------- make_icon (§38) -----------------------------
def test_make_icon_gera_ico_multi_resolucao(tmp_path):
    import make_icon
    out = tmp_path / "openrotas.ico"
    got = make_icon.gerar(str(out))
    if got is None:
        pytest.skip("Pillow indisponível neste ambiente")
    assert out.exists()
    from PIL import Image
    im = Image.open(str(out))
    assert im.format == "ICO"
    assert (256, 256) in set(im.info.get("sizes", []))   # alta resolução presente

def test_rm_reparar_tudo_sem_base_url():
    # sem base_url e com as essenciais presentes: nada a baixar, verificação OK, ok=True
    rel = rm.reparar_tudo(osrm_cfg={}, base_url="")
    assert rel["verificacao"]["ok"] is True
    assert rel["bases_atualizadas"] == [] and rel["grafo"] is None and rel["ok"] is True
    assert rel["bases_reparadas"] == []

def test_rm_reparar_corrompidos_rebaixa_forcado(tmp_path, monkeypatch):
    # Uma base corrompida na MESMA versão deve ser re-baixada (forçado) da Release e auto-curar.
    import hashlib
    monkeypatch.setattr(cfg, "user_data_dir", lambda: tmp_path / "ud")   # isola do perfil real
    monkeypatch.setattr(rm, "_manifesto_local_path", lambda: tmp_path / "ml.json")
    # arquivo "bom" servido via file:// (base_url = pasta); sha256 confere.
    srv = tmp_path / "srv"; srv.mkdir()
    bom = b"PARQUET-BOM-SIMULADO" * 100
    (srv / "snirh_rios.csv").write_bytes(bom)
    base_url = srv.as_uri()
    remoto = {"recursos": {"snirh_rios": {"versao": "2026.10", "arquivo": "snirh_rios.csv",
                                          "sha256": hashlib.sha256(bom).hexdigest()}}}
    res = rm.reparar_corrompidos(remoto, base_url, ["snirh_rios"])
    assert len(res) == 1 and res[0]["chave"] == "snirh_rios" and res[0]["ok"] is True
    # o arquivo foi gravado no override do perfil do usuário
    _, reg = rm._registry()
    assert (reg.override_dir / "snirh_rios.csv").read_bytes() == bom

def test_rm_reparar_corrompidos_pula_grafo():
    res = rm.reparar_corrompidos({"recursos": {}}, "http://x/", ["osrm_brasil"])
    assert res[0]["ok"] is False and "provision" in res[0]["detalhe"].lower()

def test_rm_reparar_tudo_grafo_nao_polui_ok(monkeypatch):
    # manifesto remoto sobe só o grafo (provisionável): NÃO deve entrar em bases_atualizadas
    # nem derrubar o ok (as bases essenciais estão presentes e legíveis).
    remoto = {"recursos": {"osrm_brasil": {"versao": "2099.1", "arquivo": "brazil-osrm-mld.tar.gz"}}}
    monkeypatch.setattr(rm, "carregar_manifesto_remoto", lambda *a, **k: remoto)
    rel = rm.reparar_tudo(osrm_cfg={}, base_url="https://exemplo/dados/")
    assert rel["bases_atualizadas"] == []          # grafo excluído do atualizar
    assert rel["ok"] is True and isinstance(rel["ok"], bool)

def test_rm_reparar_tudo_manifesto_remoto_ausente_nao_quebra(monkeypatch):
    # manifesto remoto indisponível → {} → sem atualizações, não levanta, ok continua True
    monkeypatch.setattr(rm, "carregar_manifesto_remoto", lambda *a, **k: {})
    rel = rm.reparar_tudo(osrm_cfg={}, base_url="http://exemplo.invalido/x")
    assert rel["ok"] is True and rel["bases_atualizadas"] == []

def test_osrm_telemetria_nunca_quebra():
    from engines import osrm_manager as _osrm
    # helper defensivo: não deve levantar mesmo com campos arbitrários
    _osrm._telemetria("teste", ok=True, ms=1.0)

def test_launcher_tel_nunca_quebra():
    import launcher
    # helper defensivo de telemetria do launcher: não deve levantar
    launcher._tel("teste_launcher", ok=True, ms=2.0)

def test_launcher_flags_nao_reconhecidas():
    import launcher
    # sem flag (abertura normal da janela) → nada a sinalizar
    assert launcher._flags_nao_reconhecidas([]) == []
    # a flag interna do processo-servidor é ignorada (não é "desconhecida")
    assert launcher._flags_nao_reconhecidas([launcher._SERVIDOR_FLAG]) == []
    # modo digitado errado ou flag auxiliar solta → sinalizadas (para avisar em vez de subir o app)
    assert launcher._flags_nao_reconhecidas(["--isocrona"]) == ["--isocrona"]
    assert launcher._flags_nao_reconhecidas(["--html", "--xyz"]) == ["--html", "--xyz"]


# ------------------------------- benchmark (§9/§29) --------------------------
def test_benchmark_exporta_json(tmp_path):
    from engines import benchmark as bm
    resultados = [{"rotulo": "OSRM local", "taxa_sucesso_pct": 100.0, "lat_media_s": 0.03},
                  {"rotulo": "OSRM público", "taxa_sucesso_pct": 90.0, "lat_media_s": 0.8}]
    out = tmp_path / "bench.json"
    assert bm.exportar_json(resultados, str(out)) is True
    import json as _json
    dados = _json.loads(out.read_text(encoding="utf-8"))
    assert "gerado_em" in dados and len(dados["resultados"]) == 2
    assert dados["resultados"][0]["rotulo"] == "OSRM local"

def test_benchmark_registrar_telemetria_nunca_quebra():
    from engines import benchmark as bm
    # não deve levantar mesmo com resultados mínimos/estranhos
    bm.registrar_telemetria([{"rotulo": "x", "sucessos": 1, "lat_media_s": 0.05}])
    bm.registrar_telemetria([{}])

def test_benchmark_carregar_pares_amostra():
    from engines import benchmark as bm
    pares = bm.carregar_pares(bm.SAMPLE)
    assert len(pares) >= 5 and len(pares[0]) == 5   # nome + 4 coords


# ---------------------- resource_manager: manifesto (§18/§19/§45) ------------
def test_rm_manifesto_embarcado_valido():
    m = rm.carregar_manifesto()
    assert m.get("schema") == 1
    rec = m.get("recursos", {})
    assert "municipios" in rec and "osrm_brasil" in rec
    assert all("versao" in r and "arquivo" in r for r in rec.values())

def test_rm_versao_maior_numerica():
    assert rm._versao_maior("2026.11", "2026.10") is True
    assert rm._versao_maior("2026.10", "2026.10") is False
    assert rm._versao_maior("2026.9", "2026.10") is False   # compara por campo, não texto

def test_rm_verificar_atualizacoes_detecta_versao_nova():
    local = {"recursos": {"municipios": {"versao": "2026.10", "arquivo": "municipios.parquet"}}}
    remoto = {"recursos": {"municipios": {"versao": "2026.11", "arquivo": "municipios.parquet",
                                          "sha256": "abc", "obrigatorio": True}}}
    pend = rm.verificar_atualizacoes(remoto, local)
    assert len(pend) == 1 and pend[0]["chave"] == "municipios"
    assert pend[0]["versao_remota"] == "2026.11" and pend[0]["obrigatorio"] is True

def test_rm_verificar_atualizacoes_sem_mudanca():
    mesmo = {"recursos": {"municipios": {"versao": "2026.10", "arquivo": "municipios.parquet"}}}
    assert rm.verificar_atualizacoes(mesmo, mesmo) == []

def test_rm_manifesto_local_overlay_idempotente(tmp_path, monkeypatch):
    # Após registrar uma versão instalada no manifesto local, carregar_manifesto() a sobrepõe
    # e verificar_atualizacoes deixa de reportar o recurso como pendente (§19, idempotência).
    local = tmp_path / "manifest.local.json"
    monkeypatch.setattr(rm, "_manifesto_local_path", lambda: local)
    rm._registrar_versao_local("municipios", "2026.11")
    m = rm.carregar_manifesto()
    assert m["recursos"]["municipios"]["versao"] == "2026.11"   # instalada vence a embarcada
    remoto = {"recursos": {"municipios": {"versao": "2026.11", "arquivo": "municipios.parquet"}}}
    assert rm.verificar_atualizacoes(remoto) == []              # nada pendente

def test_rm_atualizar_pula_grafo_osrm(tmp_path, monkeypatch):
    # osrm_brasil é provisionável (grafo), não base de cópia: atualizar NÃO deve baixá-lo.
    monkeypatch.setattr(rm, "_manifesto_local_path", lambda: tmp_path / "ml.json")
    remoto = {"recursos": {"osrm_brasil": {"versao": "2099.1", "arquivo": "brazil-osrm-mld.tar.gz"}}}
    res = rm.atualizar(remoto, "http://exemplo.invalido/base")
    assert len(res) == 1 and res[0]["chave"] == "osrm_brasil"
    assert res[0]["ok"] is False and "provision" in res[0]["detalhe"].lower()

def test_rm_baixar_e_verificar_rejeita_hash_errado(tmp_path):
    origem = tmp_path / "origem.bin"
    origem.write_bytes(b"conteudo-de-teste")
    destino = tmp_path / "destino.bin"
    r = rm.baixar_e_verificar(origem.as_uri(), destino, sha256="0" * 64)
    assert r["ok"] is False and not destino.exists()        # hash não bate → descartado

def test_rm_carregar_manifesto_remoto_via_file_url(tmp_path):
    import json as _json
    man = {"schema": 1, "recursos": {"municipios": {"versao": "2026.11", "arquivo": "municipios.parquet"}}}
    (tmp_path / "manifest.json").write_text(_json.dumps(man), encoding="utf-8")
    base_url = (tmp_path).as_uri()        # file:///.../  → + /manifest.json
    got = rm.carregar_manifesto_remoto(base_url)
    assert got.get("schema") == 1 and "municipios" in got.get("recursos", {})

def test_rm_carregar_manifesto_remoto_ausente_degrada():
    assert rm.carregar_manifesto_remoto("http://127.0.0.1:5999/naoexiste", timeout=1) == {}

def test_rm_baixar_e_verificar_aceita_hash_certo(tmp_path):
    import hashlib
    dados = b"conteudo-de-teste-ok"
    origem = tmp_path / "origem.bin"
    origem.write_bytes(dados)
    destino = tmp_path / "destino.bin"
    h = hashlib.sha256(dados).hexdigest()
    r = rm.baixar_e_verificar(origem.as_uri(), destino, sha256=h)
    assert r["ok"] is True and destino.exists() and destino.read_bytes() == dados


# ---------------------------- telemetry: exec_profile (§42) ------------------
def test_telemetria_registra_e_lista(tmp_path):
    arq = tmp_path / "t.jsonl"
    assert exec_profile.registrar({"evento": "rota", "ms": 12.0, "ok": True}, caminho=arq) is True
    evs = exec_profile.listar(10, caminho=arq)
    assert len(evs) == 1 and evs[0]["evento"] == "rota"
    assert "ts" in evs[0] and "iso" in evs[0]       # carimbo automático

def test_telemetria_cronometro_mede_e_registra(tmp_path):
    arq = tmp_path / "t.jsonl"
    with exec_profile.cronometro("lote", caminho=arq, motor="osrm_local"):
        pass
    evs = exec_profile.listar(10, caminho=arq)
    assert len(evs) == 1
    assert evs[0]["evento"] == "lote" and evs[0]["ok"] is True
    assert evs[0]["motor"] == "osrm_local" and isinstance(evs[0]["ms"], (int, float))

def test_telemetria_cronometro_registra_erro_e_relevanta(tmp_path):
    arq = tmp_path / "t.jsonl"
    with pytest.raises(ValueError):
        with exec_profile.cronometro("rota", caminho=arq):
            raise ValueError("falhou")
    evs = exec_profile.listar(10, caminho=arq)
    assert len(evs) == 1 and evs[0]["ok"] is False and "ValueError" in evs[0]["erro"]

def test_telemetria_resumo_agrega(tmp_path):
    arq = tmp_path / "t.jsonl"
    for ms in (10.0, 20.0, 30.0, 40.0):
        exec_profile.registrar({"evento": "rota", "ms": ms, "ok": True}, caminho=arq)
    exec_profile.registrar({"evento": "rota", "ms": 50.0, "ok": False}, caminho=arq)
    r = exec_profile.resumo(caminho=arq)
    assert r["total"] == 5 and r["ok"] == 4 and r["erro"] == 1
    assert r["por_evento"]["rota"] == 5
    t = r["tempos"]["rota"]
    assert t["n"] == 5 and t["max_ms"] == 50.0 and 10.0 <= t["media_ms"] <= 50.0

def test_telemetria_resumo_vazio(tmp_path):
    r = exec_profile.resumo(caminho=tmp_path / "inexistente.jsonl")
    assert r["total"] == 0 and r["por_evento"] == {} and r["tempos"] == {}

def test_telemetria_limpar(tmp_path):
    arq = tmp_path / "t.jsonl"
    exec_profile.registrar({"evento": "x"}, caminho=arq)
    assert arq.exists()
    assert exec_profile.limpar(caminho=arq) is True and not arq.exists()

def test_telemetria_nunca_levanta_em_caminho_invalido():
    # diretório inexistente/sem permissão não deve propagar exceção
    assert exec_profile.registrar({"evento": "x"}, caminho="/proc/openrotas-nao-existe/t.jsonl") is False


# -------------------------------- diagnostics --------------------------------
def test_diagnostico_essencial_ok(capsys, monkeypatch):
    monkeypatch.setenv("OPENROTAS_NO_NET", "1")   # não bate na rede durante o teste
    rc = diagnostics.executar(verbose=True)
    out = capsys.readouterr().out
    assert rc == 0
    assert "tudo essencial OK" in out
    assert "Offline:" in out and "Perfil de desempenho:" in out

def test_preflight_download_modulos_importam(capsys, monkeypatch):
    # Guarda anti-regressão: todos os módulos da Central de Dados precisam importar (classe
    # 'No module named osrm_manager'). Com NO_NET=1 pula o fetch e checa só os imports.
    monkeypatch.setenv("OPENROTAS_NO_NET", "1")
    rc = diagnostics.preflight_download(verbose=True)
    out = capsys.readouterr().out
    assert rc == 0, out
    assert "PREFLIGHT DO DOWNLOAD DO MAPA: OK" in out


# ----------------------------- app_update (§19) ------------------------------
def test_update_versao_atual_casa_com_config():
    assert app_update.versao_atual() == cfg.APP_VERSION

def test_update_comparacao_de_versao():
    assert app_update.ha_atualizacao("0.1.0", "0.2.0") is True
    assert app_update.ha_atualizacao("0.1.0", "0.1.0") is False
    assert app_update.ha_atualizacao("0.2.0", "0.1.9") is False   # por campo, não texto
    assert app_update.ha_atualizacao("v0.1.0", "v0.1.1") is True   # tolera prefixo 'v'

def test_update_verificar_offline_degrada(monkeypatch):
    # consulta indisponível → não quebra, marca disponivel=False
    monkeypatch.setattr(app_update, "consultar_release",
                        lambda *a, **k: {"ok": False, "versao": "", "tag": "", "url_instalador": None})
    v = app_update.verificar()
    assert v["disponivel"] is False and v["ha_atualizacao"] is False
    assert v["atual"] == cfg.APP_VERSION

def test_update_verificar_detecta_nova(monkeypatch):
    monkeypatch.setattr(app_update, "consultar_release",
                        lambda *a, **k: {"ok": True, "versao": "99.0.0",
                                         "url_instalador": "http://x/OpenRotas-Setup.exe", "tag": "app-latest"})
    v = app_update.verificar()
    assert v["disponivel"] is True and v["ha_atualizacao"] is True and v["remota"] == "99.0.0"

def test_update_baixar_sem_url_nao_quebra():
    r = app_update.baixar_instalador("")
    assert r["ok"] is False

def test_update_extrair_versao():
    assert app_update._extrair_versao("OpenRotas 0.2.0") == "0.2.0"
    assert app_update._extrair_versao("v1.10.3") == "1.10.3"
    assert app_update._extrair_versao("sem numero") == ""

def test_update_compara_0_10_maior_que_0_9():
    # bug clássico de comparação textual: 0.10 deve ser > 0.9
    assert app_update.ha_atualizacao("0.9", "0.10") is True


# --------------------- osrm parts: classificação de erro ---------------------
def test_osrm_parte_ausente_classifica():
    import urllib.error
    he404 = urllib.error.HTTPError("u", 404, "nf", {}, None)
    he500 = urllib.error.HTTPError("u", 500, "err", {}, None)
    assert osrm._parte_ausente(he404) is True
    assert osrm._parte_ausente(he500) is False          # 5xx = transitório, não fim
    assert osrm._parte_ausente(urllib.error.URLError(FileNotFoundError())) is True
    assert osrm._parte_ausente(urllib.error.URLError(ConnectionResetError())) is False


# ============================================================================
#  CONTEÚDO COMPLETO DO BRASIL — graph_url turnkey, provisionamento e Central
#  de Dados (a opção DENTRO do software para baixar/instalar tudo). §12/§32.
# ============================================================================

# ---- desktop_config: URL turnkey do grafo ----
def test_url_grafo_padrao_aponta_para_release():
    u = cfg.url_grafo_padrao()
    assert u.startswith("https://github.com/") and "osrm-brasil-latest" in u
    assert u.endswith("brazil-osrm-mld.tar.gz.part00")

def test_resolver_graph_url_configurado_vence():
    conf = {"osrm": {"graph_url": "http://meu/servidor/grafo.tar.gz.part00"}}
    assert cfg.resolver_graph_url(conf) == "http://meu/servidor/grafo.tar.gz.part00"

def test_resolver_graph_url_cai_no_padrao():
    assert cfg.resolver_graph_url({"osrm": {"mode": "docker"}}) == cfg.url_grafo_padrao()
    assert cfg.resolver_graph_url({}) == cfg.url_grafo_padrao()

def test_resolver_graph_url_env_override(monkeypatch):
    import importlib
    monkeypatch.setenv("OPENROTAS_REPO", "fulano/Fork")
    monkeypatch.setenv("OPENROTAS_GRAFO_TAG", "osrm-x")
    importlib.reload(cfg)
    try:
        u = cfg.url_grafo_padrao()
        assert "fulano/Fork" in u and "osrm-x" in u
    finally:
        monkeypatch.delenv("OPENROTAS_REPO", raising=False)
        monkeypatch.delenv("OPENROTAS_GRAFO_TAG", raising=False)
        importlib.reload(cfg)   # restaura os padrões para os demais testes


# ---- osrm_manager: callback de progresso no download em partes ----
def test_baixar_partes_emite_progresso(tmp_path):
    srv = tmp_path / "srv"; srv.mkdir()
    (srv / "g.tar.gz.part00").write_bytes(b"A" * (1024 * 300))
    (srv / "g.tar.gz.part01").write_bytes(b"B" * (1024 * 300))
    base = (srv / "g.tar.gz.part").as_uri()
    eventos = []
    ok = osrm._baixar_partes(base, tmp_path / "out.tar.gz", progresso=lambda ev: eventos.append(ev))
    assert ok is True and (tmp_path / "out.tar.gz").exists()
    assert any(e.get("fase") == "baixando" for e in eventos)
    assert any(e.get("parte_concluida") for e in eventos)   # fim de cada parte sinalizado

def test_baixar_partes_progresso_que_levanta_nao_quebra(tmp_path):
    srv = tmp_path / "srv"; srv.mkdir()
    (srv / "g.tar.gz.part00").write_bytes(b"A" * 2048)
    base = (srv / "g.tar.gz.part").as_uri()
    def _boom(ev):
        raise RuntimeError("callback ruim")
    # um progresso defeituoso NUNCA pode abortar/corromper o download
    assert osrm._baixar_partes(base, tmp_path / "out.tar.gz", progresso=_boom) is True

def test_garantir_grafo_emite_pronto_quando_ja_existe(tmp_path):
    destino = tmp_path / "dl"; destino.mkdir()
    (destino / "brazil-latest.osrm").write_bytes(b"x")
    eventos = []
    got = osrm.garantir_grafo({}, destino, progresso=lambda ev: eventos.append(ev))
    assert got == str(destino / "brazil-latest.osrm")
    assert any(e.get("fase") == "pronto" for e in eventos)


# ---- provisionamento: camada pura ----
def test_prov_pasta_de_dados_existe():
    p = prov.pasta_de_dados()
    assert p.name == "data_local"

def test_prov_humano_bytes():
    assert prov.humano_bytes(0) == "0 B"
    assert prov.humano_bytes(1536).endswith("KB")
    assert prov.humano_bytes(6_700_000_000).endswith("GB")
    assert prov.humano_bytes("xx") == "—"

def test_prov_inventario_estrutura():
    inv = prov.inventario()
    assert "recursos" in inv and "resumo" in inv and "offline" in inv
    assert inv["resumo"]["total"] >= 1
    # as bases essenciais estão instaladas neste repo (reassembladas) → nenhum obrigatório falta
    assert isinstance(inv["resumo"]["faltam_obrigatorios"], list)
    # osrm_brasil é opcional: aparece em faltam_opcionais sse e só se o grafo não está instalado.
    # Invariante portável (alguns ambientes já trazem o grafo).
    assert ("osrm_brasil" in inv["resumo"]["faltam_opcionais"]) == (not prov.grafo_instalado())

def test_prov_grafo_instalado_reflete_a_pasta():
    # grafo_instalado() é um bool coerente com a pasta de dados real: True quando há a família MLD
    # essencial completa (ou um arquivo-base avulso), False caso contrário. A cobertura exata da
    # detecção (família incompleta, base avulso) está nos testes isolados com tmp_path abaixo; aqui
    # só garantimos que no ambiente real a resposta é booleana e não levanta. Portável.
    assert isinstance(prov.grafo_instalado(), bool)

def test_prov_grafo_instalado_exige_familia_mld_completa(tmp_path, monkeypatch):
    # Robustez: 'instalado' só é True com um grafo MLD COMPLETO. Uma extração PARCIAL não pode
    # travar o botão 'Baixar o mapa' nem ligar o roteamento sobre um grafo quebrado.
    d = tmp_path / "dados"; d.mkdir()
    monkeypatch.setattr(prov, "pasta_de_dados", lambda: d)
    # família incompleta (falta .cell_metrics) → NÃO instalado
    for suf in (".mldgr", ".partition", ".cells"):
        (d / ("brazil-latest.osrm" + suf)).write_bytes(b"x")
    assert prov.grafo_instalado() is False
    # completa os essenciais do MLD → instalado
    (d / "brazil-latest.osrm.cell_metrics").write_bytes(b"x")
    assert prov.grafo_instalado() is True
    assert not (d / "brazil-latest.osrm").exists()        # sem arquivo-base literal, mesmo assim OK

def test_prov_grafo_instalado_arquivo_base_avulso(tmp_path, monkeypatch):
    # Compat: um '<x>.osrm' avulso (formatos antigos) conta como instalado.
    d = tmp_path / "dados"; d.mkdir()
    monkeypatch.setattr(prov, "pasta_de_dados", lambda: d)
    (d / "brazil-latest.osrm").write_bytes(b"x")
    assert prov.grafo_instalado() is True

def _carregar_scanner_segredos():
    import importlib.util, pathlib
    p = pathlib.Path(__file__).resolve().parents[1] / "installer" / "verificar_segredos.py"
    spec = importlib.util.spec_from_file_location("verificar_segredos", p)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

def _jwt_fake(role):
    import base64 as _b64
    h = _b64.urlsafe_b64encode(b'{"alg":"HS256","typ":"JWT"}').rstrip(b"=").decode()
    pl = _b64.urlsafe_b64encode(('{"role":"%s","iss":"supabase"}' % role).encode()).rstrip(b"=").decode()
    return h + "." + pl + ".assinaturafake123"

def test_scanner_segredos_detecta_chave_real(tmp_path):
    seg = _carregar_scanner_segredos()
    # (1) JWT service_role e (2) sb_secret_<material> → DETECTADOS.
    (tmp_path / "vaza1.json").write_text('{"SUPABASE_KEY": "%s"}' % _jwt_fake("service_role"),
                                         encoding="utf-8")
    (tmp_path / "vaza2.env").write_text("SB=" + "sb" + "_secret_" + "ABCDEFGHIJKLMNOP0123456789",
                                        encoding="utf-8")
    achados = seg.escanear_segredos([str(tmp_path)])
    marcs = {a["marcador"] for a in achados}
    assert "service_role JWT" in marcs and "sb_secret key" in marcs
    assert seg.main([str(tmp_path)]) == 2

def test_scanner_segredos_ignora_mencao_defensiva(tmp_path):
    seg = _carregar_scanner_segredos()
    # A PALAVRA em código de validação + um JWT 'anon' (role público) → NÃO é achado.
    (tmp_path / "guard.py").write_text(
        '# nunca embutir service_role no cliente\n'
        'if "service_role" in chave: raise SystemExit("remova a service_role")\n'
        'ANON = "%s"\n' % _jwt_fake("anon"), encoding="utf-8")
    assert seg.escanear_segredos([str(tmp_path)]) == []
    assert seg.main([str(tmp_path)]) == 0

def test_scanner_segredos_ignora_dir_de_testes(tmp_path):
    seg = _carregar_scanner_segredos()
    d = tmp_path / "tests"; d.mkdir()
    (d / "fixture.json").write_text('{"k":"%s"}' % _jwt_fake("service_role"), encoding="utf-8")
    assert seg.escanear_segredos([str(tmp_path)]) == []   # dir 'tests' é ignorado

def test_prov_bases_ausentes_usa_registry(monkeypatch):
    # Regressão do bug _registry() inexistente (bases_ausentes caía sempre no except → []).
    # Com um registry que reporta "nada instalado", deve LISTAR as bases não-osrm faltando.
    import local_data
    class _RegVazio:
        def existe(self, chave): return False
    monkeypatch.setattr(prov, "_registry", lambda: (local_data, _RegVazio()))
    faltam = prov.bases_ausentes()
    nao_osrm = [d.chave for d in local_data.CATALOGO if d.formato != "osrm"]
    assert faltam == nao_osrm and len(faltam) > 0              # agora lista de verdade
    assert "osrm_brasil" not in faltam                         # o grafo é tratado à parte
    # e _registry() de verdade existe e devolve (módulo, registry) — não mais NameError
    _m, _r = prov._registry()
    assert _m is local_data and hasattr(_r, "existe")

def test_prov_checar_espaco_para_grafo(monkeypatch):
    import shutil as _sh
    from collections import namedtuple
    DU = namedtuple("DU", "total used free")
    # Disco cheio → bloqueia com mensagem clara.
    monkeypatch.setattr(prov.shutil, "disk_usage", lambda *_a: DU(0, 0, 1 * 1024 ** 3))
    r = prov.checar_espaco_para_grafo()
    assert r["ok"] is False and "insuficiente" in r["detalhe"]
    # Disco folgado → libera.
    monkeypatch.setattr(prov.shutil, "disk_usage", lambda *_a: DU(0, 0, 50 * 1024 ** 3))
    assert prov.checar_espaco_para_grafo()["ok"] is True
    # Não mediu (exceção) → NÃO bloqueia (evita falso negativo).
    def _boom(*_a): raise OSError("no disk")
    monkeypatch.setattr(prov.shutil, "disk_usage", _boom)
    assert prov.checar_espaco_para_grafo()["ok"] is True

def test_prov_verificar_integridade_detecta_corrompida(monkeypatch, tmp_path):
    # Uma base essencial "corrompida" (existe mas não abre) deve ser detectada.
    monkeypatch.setattr(prov, "pasta_de_dados", lambda: tmp_path / "data_local")
    import desktop_config as _cfg
    monkeypatch.setattr(_cfg, "user_data_dir", lambda: tmp_path)
    import local_data
    _mod, reg = prov._registry()
    # força: tudo "existe", mas 'municipios' não é legível (corrompida)
    monkeypatch.setattr(type(reg), "existe", lambda self, chave: True)
    monkeypatch.setattr(type(reg), "legivel", lambda self, chave: chave != "municipios")
    monkeypatch.setattr(prov, "_registry", lambda: (local_data, reg))
    r = prov.verificar_integridade()
    assert r["ok"] is False and "municipios" in r["corrompidas"]

def test_prov_pasta_por_tipo(monkeypatch, tmp_path):
    import desktop_config as _cfg
    monkeypatch.setattr(_cfg, "user_data_dir", lambda: tmp_path)
    assert prov.pasta_por_tipo("exports") == tmp_path / "exports"
    assert prov.pasta_por_tipo("logs") == tmp_path / "logs"
    assert str(prov.pasta_por_tipo("downloads")).endswith("Downloads")

def test_prov_checar_atualizacao_defensivo(monkeypatch):
    # Com app_update dizendo que há nova versão, o wrapper repassa fielmente.
    import app_update as _au
    monkeypatch.setattr(_au, "verificar", lambda *a, **k: {
        "atual": "0.2.0", "remota": "0.3.0", "ha_atualizacao": True,
        "url_instalador": "https://x/OpenRotas.Setup.exe", "disponivel": True, "detalhe": "nova"})
    r = prov.checar_atualizacao()
    assert r["ha_atualizacao"] is True and r["url_instalador"].endswith(".exe")
    # Se app_update levanta, o wrapper não propaga (disponivel=False).
    monkeypatch.setattr(_au, "verificar", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
    r2 = prov.checar_atualizacao()
    assert r2["disponivel"] is False and r2["ha_atualizacao"] is False

def test_dossie_lote_progresso(monkeypatch, tmp_path):
    # Novo callback de progresso do lote: emite feito/total por par (motor mockado, sem rede).
    import sys as _sys, pathlib as _pl
    _sys.path.insert(0, str(_pl.Path(__file__).resolve().parents[1]))
    from geo import dossie_lote, dossie_rota
    monkeypatch.setattr(dossie_lote, "ler_pares",
                        lambda _c: [("-46.6,-23.5", "-43.2,-22.9"), ("-47.9,-15.8", "-49.2,-16.7")])
    monkeypatch.setattr(dossie_rota, "dossie", lambda pts, **k: {"distancia_km": 1.0})
    eventos = []
    r = dossie_lote.processar_lote("x.csv", caminho_out=str(tmp_path / "out.xlsx"),
                                   progresso=lambda ev: eventos.append(ev))
    assert [e["feito"] for e in eventos] == [1, 2] and eventos[0]["total"] == 2
    assert r["linhas"] == 2                       # 2 pares processados (grava depende de openpyxl)

def test_prov_processar_dossie_lote_grava_em_exports(monkeypatch, tmp_path):
    import sys as _sys, pathlib as _pl
    _sys.path.insert(0, str(_pl.Path(__file__).resolve().parents[1]))
    from geo import dossie_lote
    monkeypatch.setattr(prov, "pasta_por_tipo", lambda tipo: tmp_path / tipo)
    capt = {}
    def _fake(cin, caminho_out=None, usar_osrm=False, progresso=None, **k):
        capt["out"] = caminho_out
        return {"ok": True, "linhas": 3, "erros": 0, "saida": caminho_out}
    monkeypatch.setattr(dossie_lote, "processar_lote", _fake)
    r = prov.processar_dossie_lote("planilha.csv")
    assert r["ok"] is True and r["linhas"] == 3
    assert "exports" in capt["out"] and capt["out"].endswith(".xlsx")

def test_prov_baixar_atualizacao(monkeypatch):
    import app_update as _au
    # há atualização → baixa e repassa o caminho
    monkeypatch.setattr(prov, "checar_atualizacao", lambda: {
        "disponivel": True, "ha_atualizacao": True,
        "url_instalador": "https://x/OpenRotas.Setup.exe", "remota": "0.3.0"})
    monkeypatch.setattr(_au, "baixar_instalador",
                        lambda url, *a, **k: {"ok": True, "caminho": r"C:\up\OpenRotas.Setup.exe",
                                             "detalhe": "ok"})
    r = prov.baixar_atualizacao()
    assert r["ok"] is True and r["caminho"].endswith("OpenRotas.Setup.exe") and r["versao"] == "0.3.0"
    # sem atualização → não baixa
    monkeypatch.setattr(prov, "checar_atualizacao", lambda: {
        "disponivel": True, "ha_atualizacao": False, "url_instalador": None})
    assert prov.baixar_atualizacao()["ok"] is False

def test_prov_resumo_desempenho_sem_dados(monkeypatch):
    import exec_profile as _ep
    monkeypatch.setattr(_ep, "resumo", lambda *a, **k: {"total": 0, "ok": 0, "erro": 0, "tempos": {}})
    r = prov.resumo_desempenho()
    assert r["total"] == 0 and "texto" in r and isinstance(r["texto"], str)

def test_prov_resumo_desempenho_com_rota(monkeypatch):
    import exec_profile as _ep
    monkeypatch.setattr(_ep, "resumo", lambda *a, **k: {
        "total": 10, "ok": 9, "erro": 1,
        "tempos": {"rota": {"n": 8, "media_ms": 42, "p95_ms": 99}}})
    r = prov.resumo_desempenho()
    assert "média 42ms" in r["texto"] and "p95 99ms" in r["texto"]

def test_prov_relatorio_diagnostico_e_texto():
    txt = prov.relatorio_diagnostico()
    assert isinstance(txt, str) and "OpenRotas" in txt and "Diagnóstico" in txt

def test_prov_saude_nacional_texto(monkeypatch):
    # Não depende de bases reais: monkeypatcha o módulo de auditoria e confirma o repasse.
    import sys as _sys, pathlib as _pl
    _sys.path.insert(0, str(_pl.Path(__file__).resolve().parents[1] / "audit"))
    import saude_nacional as _sn
    monkeypatch.setattr(_sn, "diagnostico", lambda rapido=True: {"itens": [], "ok": True})
    monkeypatch.setattr(_sn, "render_texto", lambda d: "SAÚDE NACIONAL: ok")
    assert prov.saude_nacional_texto() == "SAÚDE NACIONAL: ok"
    # Se o módulo levanta, o wrapper devolve mensagem legível (não propaga).
    monkeypatch.setattr(_sn, "diagnostico", lambda rapido=True: (_ for _ in ()).throw(RuntimeError("x")))
    assert "Falha ao avaliar" in prov.saude_nacional_texto()

def test_launcher_silenciar_logs_sensiveis():
    import logging as _lg
    import launcher
    # zera para um nível baixo e confirma que a higiene sobe para WARNING
    for nome in ("httpx", "urllib3", "httpcore"):
        _lg.getLogger(nome).setLevel(_lg.INFO)
    launcher._silenciar_logs_sensiveis()
    for nome in ("httpx", "urllib3", "httpcore"):
        assert _lg.getLogger(nome).level == _lg.WARNING

def _patch_health(monkeypatch, valor):
    # status_motor_rotas resolve o osrm_manager por 'import osrm_manager' (plano); o teste importa
    # 'engines.osrm_manager' — podem ser objetos distintos. Patcha health em AMBOS, defensivamente.
    import sys as _sys
    monkeypatch.setattr(osrm, "health", lambda *a, **k: valor)
    try:
        import osrm_manager as _flat
        monkeypatch.setattr(_flat, "health", lambda *a, **k: valor)
    except Exception:
        pass
    if "osrm_manager" in _sys.modules:
        monkeypatch.setattr(_sys.modules["osrm_manager"], "health", lambda *a, **k: valor, raising=False)

def test_prov_status_motor_rotas_local_no_ar(monkeypatch):
    # Motor local respondendo ao health → status 'no ar', NÃO usando público.
    _patch_health(monkeypatch, True)
    st = prov.status_motor_rotas(conf={"osrm": {"mode": "auto", "port": 5000}})
    assert st["local_no_ar"] is True and st["usando_publico"] is False
    assert "NO AR" in st["detalhe"] and "127.0.0.1:5000" in st["url_local"]
    # Com o motor no ar, as ANÁLISES usam distância rodoviária real (não estimativa).
    assert "rodoviário" in st["ferramentas_modo"].lower()

def test_prov_status_motor_rotas_cai_no_publico(monkeypatch):
    # Motor NÃO responde → usando público; a mensagem explica (não é falha silenciosa).
    _patch_health(monkeypatch, False)
    st_off = prov.status_motor_rotas(conf={"osrm": {"mode": "off"}})
    assert st_off["local_no_ar"] is False and st_off["usando_publico"] is True
    assert "desligado" in st_off["detalhe"]
    assert "estim" in st_off["ferramentas_modo"].lower()
    st_on = prov.status_motor_rotas(conf={"osrm": {"mode": "auto"}})
    assert st_on["usando_publico"] is True and "público" in st_on["detalhe"]
    # Sem motor local no ar, as ferramentas do desktop caem na ESTIMATIVA (não no OSRM público).
    assert "estim" in st_on["ferramentas_modo"].lower()

def test_prov_roteamento_local_ativo(tmp_path):
    assert prov.roteamento_local_ativo(conf={"osrm": {"mode": "off"}}) is False
    assert prov.roteamento_local_ativo(conf={"osrm": {}}) is False
    assert prov.roteamento_local_ativo(conf={"osrm": {"mode": "auto"}}) is True
    assert prov.roteamento_local_ativo(conf={"osrm": {"mode": "local"}}) is True

def test_prov_baixar_grafo_em_partes_para_a_pasta(tmp_path, monkeypatch):
    import io, tarfile, gzip
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w") as t:
        dado = b"OSRM" * 4000
        info = tarfile.TarInfo("brazil-latest.osrm"); info.size = len(dado)
        t.addfile(info, io.BytesIO(dado))
    targz = gzip.compress(buf.getvalue())
    srv = tmp_path / "srv"; srv.mkdir()
    meio = len(targz) // 2
    (srv / "brazil-osrm-mld.tar.gz.part00").write_bytes(targz[:meio])
    (srv / "brazil-osrm-mld.tar.gz.part01").write_bytes(targz[meio:])
    base = (srv / "brazil-osrm-mld.tar.gz.part").as_uri()
    destino = tmp_path / "pasta_dados"
    monkeypatch.setattr(prov, "pasta_de_dados", lambda: destino)
    eventos = []
    r = prov.baixar_grafo(progresso=lambda ev: eventos.append(ev),
                          conf={"osrm": {"graph_url": base}})
    assert r["ok"] is True and (destino / "brazil-latest.osrm").exists()
    assert any(e.get("fase") in ("baixando", "extraindo") for e in eventos)

def test_prov_baixar_grafo_familia_mld_reporta_sucesso(tmp_path, monkeypatch):
    # REGRESSÃO do "[causa: sem detalhe]": com o grafo real (OSRM 5.x MLD), garantir_grafo
    # devolve o CAMINHO-BASE 'brazil-latest.osrm' que NÃO existe como arquivo (só a família
    # '.osrm.*'). O antigo gate os.path.exists(caminho) dava False → download bem-sucedido
    # era reportado como falha "sem detalhe". Agora baixar_grafo deve reportar ok=True.
    import io, tarfile, gzip
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w") as t:
        for suf in (".cells", ".mldgr", ".partition", ".properties"):
            nome = "brazil-latest.osrm" + suf
            dado = (nome + "-data").encode() * 50
            info = tarfile.TarInfo(nome); info.size = len(dado)
            t.addfile(info, io.BytesIO(dado))
    targz = gzip.compress(buf.getvalue())
    srv = tmp_path / "srv"; srv.mkdir()
    meio = len(targz) // 2
    (srv / "brazil-osrm-mld.tar.gz.part00").write_bytes(targz[:meio])
    (srv / "brazil-osrm-mld.tar.gz.part01").write_bytes(targz[meio:])
    base = (srv / "brazil-osrm-mld.tar.gz.part").as_uri()
    destino = tmp_path / "pasta_dados"
    monkeypatch.setattr(prov, "pasta_de_dados", lambda: destino)
    r = prov.baixar_grafo(conf={"osrm": {"graph_url": base}})
    assert r["ok"] is True                                   # NÃO "sem detalhe"
    assert r["caminho"] == str(destino / "brazil-latest.osrm")
    assert not (destino / "brazil-latest.osrm").exists()     # base não existe como arquivo
    assert (destino / "brazil-latest.osrm.mldgr").exists()   # a família está instalada

def test_prov_grafo_valido_reconhece_familia(tmp_path):
    assert prov._grafo_valido("") is False
    assert prov._grafo_valido(str(tmp_path / "brazil-latest.osrm")) is False   # nada ainda
    (tmp_path / "brazil-latest.osrm.cells").write_bytes(b"x")
    assert prov._grafo_valido(str(tmp_path / "brazil-latest.osrm")) is True    # família conta
    avulso = tmp_path / "g.osrm"; avulso.write_bytes(b"x")
    assert prov._grafo_valido(str(avulso)) is True                             # avulso também

def test_prov_baixar_grafo_sem_url_nao_quebra(monkeypatch):
    # força graph_url vazio e sem padrão → mensagem clara, sem levantar
    monkeypatch.setattr(cfg, "resolver_graph_url", lambda *a, **k: "")
    r = prov.baixar_grafo(conf={"osrm": {}})
    assert r["ok"] is False and "nada a baixar" in r["detalhe"]

def test_prov_abrir_pasta_headless_e_noop(monkeypatch):
    monkeypatch.setenv("OPENROTAS_NO_NET", "1")
    assert prov.abrir_pasta_dados() is False      # em CI/headless não tenta abrir GUI

def test_prov_ativar_roteamento_local_grava_config(tmp_path, monkeypatch):
    monkeypatch.setattr(cfg, "user_data_dir", lambda: tmp_path / "ud")
    r = prov.ativar_roteamento_local()
    assert r["ok"] is True
    assert "nativo" in r and "docker" in r      # relata a via disponível (nativo embarcado / Docker)
    import json as _json
    conf = _json.loads((tmp_path / "ud" / "config" / "desktop.json").read_text(encoding="utf-8"))
    assert conf["osrm"]["mode"] == "auto"       # turnkey: nativo s/ Docker > Docker > público
    assert conf["osrm"]["graph_url"].endswith("part00")   # garante uma URL utilizável

def test_prov_ativar_roteamento_local_preserva_config(tmp_path, monkeypatch):
    import json as _json
    ud = tmp_path / "ud"; (ud / "config").mkdir(parents=True)
    (ud / "config" / "desktop.json").write_text(
        _json.dumps({"SUPABASE_URL": "https://x", "osrm": {"graph_url": "http://meu/g.part00"}}),
        encoding="utf-8")
    monkeypatch.setattr(cfg, "user_data_dir", lambda: ud)
    prov.ativar_roteamento_local()
    conf = _json.loads((ud / "config" / "desktop.json").read_text(encoding="utf-8"))
    assert conf["SUPABASE_URL"] == "https://x"               # resto preservado
    assert conf["osrm"]["mode"] == "auto"
    assert conf["osrm"]["graph_url"] == "http://meu/g.part00"  # url do usuário preservada

def test_prov_reparar_bases_emite_e_nao_quebra():
    eventos = []
    rel = prov.reparar_bases(progresso=lambda ev: eventos.append(ev))
    assert isinstance(rel, dict)
    assert any(e.get("fase") == "concluido" for e in eventos)


# ---- central_dados: fallback de texto (sem GUI) ----
def test_central_resumo_texto_tem_pasta_e_grafo():
    txt = central_dados._resumo_texto()
    assert "Central de Dados" in txt and "grafo" in txt.lower()
    assert "Pasta de dados:" in txt

def test_central_executar_texto_sem_baixar_retorna_0(capsys):
    rc = central_dados._executar_texto(baixar=False)
    assert rc == 0
    assert "--central --baixar" in capsys.readouterr().out

def test_central_main_cai_para_texto_sem_display(monkeypatch, capsys):
    # sem Tkinter utilizável, main() usa o modo texto e não levanta
    monkeypatch.setattr(central_dados, "_tkinter_disponivel", lambda: False)
    rc = central_dados.main([])
    assert rc == 0 and "Central de Dados" in capsys.readouterr().out


# ---- manifesto: completude (todas as camadas nacionais embarcadas) ----
def test_manifesto_cobre_catalogo_completo():
    m = rm.carregar_manifesto()
    rec = m.get("recursos", {})
    for chave in ("drenagem", "rodovias", "ferrovias", "pontes", "travessias", "hidrovias",
                  "eclusas", "atracadouros_terminal", "complexos_portuarios", "sinalizacao"):
        assert chave in rec, "manifesto deveria listar %r" % chave
    assert rec["drenagem"]["obrigatorio"] is True and rec["rodovias"]["obrigatorio"] is True

def test_painel_tem_bloco_dados_completos():
    # A seção "Dados completos (Brasil inteiro)" existe sempre. Seu CONTEÚDO é honesto em qualquer
    # ambiente: sem o grafo instalado, orienta baixá-lo pela Central de Dados; com o grafo presente,
    # confirma a instalação. Testar a invariante (uma das duas formas) torna o teste portável —
    # passa com ou sem o grafo no ambiente — sem afrouxar o que importa (o bloco é honesto).
    h = painel.construir_html()
    assert "Dados completos" in h
    assert ("Central de Dados" in h) or ("Grafo de roteamento do Brasil instalado" in h)


# ============================================================================
#  AUDITOR DE COBERTURA NACIONAL (§3/§4/§40/§42/§64/§65/§66) — prova objetiva,
#  números derivados dos dados REAIS instalados (nunca fabricados).
# ============================================================================

def test_uf_do_geocodigo_mapeia_prefixo():
    import pandas as pd
    s = pd.Series(["3550308", "1100015", "5300108", "2927408", "abc", "9999999"])
    uf = list(ca._uf_do_geocodigo(s))
    assert uf[0] == "SP" and uf[1] == "RO" and uf[2] == "DF" and uf[3] == "BA"
    assert pd.isna(uf[4]) and pd.isna(uf[5])      # lixo/código inexistente → NaN (sem inventar)

def test_auditoria_estados_27_e_municipios_reais():
    reg = ca._registry()
    em = ca.cobertura_estados_municipios(reg)
    assert em["estados"]["encontrado"] == 27 and em["estados"]["status"] == "OK"
    assert set(em["estados"]["presentes"]) == ca.UFS_ESPERADAS
    # número REAL de municípios (IBGE 5.570 + Fernando de Noronha) — não um valor inventado
    assert em["municipios"]["encontrado"] >= ca.MUNICIPIOS_OFICIAIS
    assert len(em["municipios"]["por_uf"]) == 27

def test_auditoria_completa_estrutura_e_camadas():
    laudo = ca.auditar(rapido=True)
    assert laudo["estados"]["encontrado"] == 27
    assert "rodovias" in laudo["camadas"] and "drenagem" in laudo["camadas"]
    rod = laudo["camadas"]["rodovias"]
    assert rod["instalado"] is True and rod["feicoes"] > 100000
    assert len(rod["ufs_presentes"]) >= 20          # malha nacional presente em (quase) todas as UFs
    assert rod["status"] in ("OK", "PARCIAL")
    r = laudo["resumo"]
    assert r["dimensoes"] == r["ok"] + r["parciais"] + r["ausentes"]

def test_auditoria_grafo_honesto():
    # Honestidade do laudo do grafo, em QUALQUER ambiente: o status é AUSENTE se e só se o grafo
    # não está instalado; se estiver instalado, o status é OK/PARCIAL — nunca mascara o estado real
    # como o oposto. (Testar a invariante torna o teste portável: alguns ambientes de CI já trazem o
    # grafo; o que importa é que o laudo diga a verdade sobre o estado que encontrou.)
    laudo = ca.auditar(rapido=True)
    g = laudo["grafo"]
    if g["instalado"]:
        assert g["status"] in ("OK", "PARCIAL")
    else:
        assert g["status"] == "AUSENTE"

def test_auditoria_presenca_uf_camada_ausente(tmp_path):
    # registry isolado sem bases → camada não instalada vira AUSENTE, sem levantar
    import local_data
    reg = local_data.LocalDataRegistry(cfg.app_root(), tmp_path / "data_local")
    r = ca.presenca_espacial_uf(reg, "eclusas", {}, amostra=None)
    # eclusas é embarcada (existe no bundle), mas sem bboxes a presença fica vazia, sem quebrar
    assert isinstance(r["ufs_presentes"], list)

def test_render_texto_tem_cabecalho_e_barra():
    laudo = ca.auditar(rapido=True)
    txt = ca.render_texto(laudo)
    assert "AUDITORIA NACIONAL" in txt and "Estados (UFs)" in txt
    assert "Municípios" in txt and "█" in txt       # barra de progresso ASCII

def test_relatorio_md_gera_arquivo(tmp_path):
    laudo = ca.auditar(rapido=True)
    out = tmp_path / "rel.md"
    got = ca.gerar_relatorio_md(laudo, out)
    assert got == str(out) and out.exists()
    conteudo = out.read_text(encoding="utf-8")
    assert conteudo.startswith("# Relatório de Cobertura Nacional")
    assert "Municípios por UF" in conteudo and "Nenhum dado foi" in conteudo  # nota de honestidade
    assert "| SP |" in conteudo                      # tabela por UF preenchida com dado real

def test_resumo_curto_e_cli():
    laudo = ca.auditar(rapido=True)
    s = ca.resumo_curto(laudo)
    assert "UFs 27/27" in s
    assert ca._cli(["--rapido"]) == 0                # essenciais presentes → saída 0

def test_painel_tem_bloco_cobertura():
    h = painel.construir_html()
    assert "Auditoria de Cobertura Nacional" in h and "Estados (UFs)" in h


# ============================================================================
#  GeoIntelligenceRepository + multimodal/fluvial (§6/§8/§24/§25) — consultas
#  espaciais unificadas (nível bbox, numpy), sobre os dados REAIS instalados.
# ============================================================================
@pytest.fixture(scope="module")
def repo():
    return georepo.GeoIntelligenceRepository()

def test_repo_arrays_projecao_leve(repo):
    a = repo.arrays("rodovias")
    assert a and a["n"] > 100000
    assert set(("lon", "lat", "xmin", "ymin", "xmax", "ymax")).issubset(a.keys())

def test_repo_camada_inexistente_degrada(repo):
    assert repo.arrays("nao_existe_xyz") is None
    assert repo.disponivel("osrm_brasil") is False      # provisionável, ausente aqui

def test_repo_intersecta_bbox_sp(repo):
    import numpy as np
    # janela sobre a cidade de São Paulo deve conter feições rodoviárias
    idx = repo.intersecta_bbox("rodovias", -46.8, -23.8, -46.3, -23.4)
    assert isinstance(idx, np.ndarray) and idx.size > 0

def test_repo_intersecta_bbox_fora_do_brasil_vazio(repo):
    # no meio do Atlântico não há feições rodoviárias
    assert repo.intersecta_bbox("rodovias", 10.0, 10.0, 11.0, 11.0).size == 0

def test_repo_proximidade_ponto(repo):
    r = repo.proximidade("drenagem", -60.02, -3.1, raio_km=25)   # Manaus/AM (muitos rios)
    assert r["total"] > 0 and r["raio_km"] == 25

def test_repo_contar_na_rota_sp_rio(repo):
    coords = [(-46.63, -23.55), (-43.20, -22.90)]
    res = repo.contar_na_rota(coords, camadas=["drenagem", "rodovias", "pontes"])
    assert res["drenagem"]["feicoes"] > 0 and res["drenagem"]["disponivel"] is True
    assert res["rodovias"]["feicoes"] > 0

def test_repo_contar_na_rota_camada_ausente(repo):
    res = repo.contar_na_rota([(-46.6, -23.5), (-46.5, -23.4)], camadas=["osrm_brasil"])
    assert res["osrm_brasil"] == {"feicoes": 0, "disponivel": False}

def test_repo_inventario_multimodal(repo):
    inv = repo.inventario_multimodal()
    assert inv["camadas"]["rodovias"]["instalado"] is True
    assert inv["camadas"]["drenagem"]["classe"] == "fluvial"
    assert "rodoviario" in inv["classes_presentes"] and "fluvial" in inv["classes_presentes"]
    assert inv["total_feicoes"] > 2_000_000

def test_repo_analise_multimodal_rota(repo):
    res = repo.analise_multimodal_rota([(-46.63, -23.55), (-43.20, -22.90)])
    assert res["cruza_rio"] is True
    assert res["por_classe"]["rodoviario"] > 0 and res["por_classe"]["fluvial"] > 0
    assert "aproximada" in res["nota"]

def test_repo_render_e_cli():
    inv = georepo.GeoIntelligenceRepository().inventario_multimodal()
    txt = georepo.render_inventario(inv)
    assert "INVENTÁRIO MULTIMODAL" in txt and "rodovias" in txt
    assert georepo._cli([]) == 0


# ============================================================================
#  TESTE NACIONAL DE ROTAS / COBERTURA DISTRIBUÍDA (§42/§43/§44/§45) — corredores
#  inter-UF nas 5 regiões, sobre os dados REAIS, pegando buracos regionais.
# ============================================================================
def test_rotas_corredores_cobrem_as_regioes():
    from audit import route_coverage as rc
    regioes = " ".join(r for *_, r in rc.CORREDORES)
    assert "Norte" in regioes and "Nordeste" in regioes and "Sul" in regioes
    assert len(rc.CORREDORES) >= 10

def test_rotas_nacional_todos_corredores_com_malha():
    from audit import route_coverage as rc
    aud = rc.auditar_rotas()
    assert aud["resumo"]["total"] == len(rc.CORREDORES)
    # toda a malha nacional deve aparecer em TODOS os corredores inter-UF (sem buraco regional)
    assert aud["resumo"]["status"] == "OK" and aud["resumo"]["ok"] == aud["resumo"]["total"]
    for c in aud["corredores"]:
        assert c["ok"] is True and c["rodoviario"] > 0

def test_rotas_amazonia_tem_rios():
    # corredores amazônicos devem cruzar muitos rios (sanidade regional — §45)
    from audit import route_coverage as rc
    aud = rc.auditar_rotas()
    amaz = [c for c in aud["corredores"] if c["corredor"].startswith(("AC→AM", "RR→PA"))]
    assert amaz and all(c["cruza_rio"] and c["fluvial"] > 0 for c in amaz)

def test_rotas_render_e_cli():
    from audit import route_coverage as rc
    aud = rc.auditar_rotas()
    txt = rc.render_texto(aud)
    assert "TESTE NACIONAL DE ROTAS" in txt and "corredores" in txt
    assert rc._cli() == 0


# ============================================================================
#  CATÁLOGO DE FONTES / RASTREABILIDADE (§14/§15/§33/§69/§70) — procedência REAL
#  lida dos dados + metadados declarados (organização/licença/portal).
# ============================================================================
def test_catalogo_estrutura_e_procedencia_real():
    from catalog import fontes
    cat = fontes.catalogo(medir=True)
    por = {r["chave"]: r for r in cat}
    assert por["municipios"]["organizacao"] == "IBGE"
    assert por["municipios"]["instalado"] and por["municipios"]["registros"] == 5571
    # procedência lida do próprio dado (não declarada): BC250, versão 2025
    assert "BC250" in (por["municipios"]["fonte_base"] or [])
    assert por["drenagem"]["registros"] > 2_000_000
    assert por["snirh_rios"]["organizacao"].startswith("ANA")

def test_catalogo_medir_false_e_rapido():
    from catalog import fontes
    cat = fontes.catalogo(medir=False)
    # sem medir: ainda traz metadados e existência, mas não procedência pesada
    assert all("organizacao" in r for r in cat)
    assert all("fonte_base" not in r for r in cat)

def test_catalogo_tipos_de_fonte_classificados():
    from catalog import fontes
    por = {r["chave"]: r for r in fontes.catalogo(medir=False)}
    assert por["municipios"]["tipo_fonte"] == "primaria"
    assert por["snirh_rios"]["tipo_fonte"] == "complementar"
    assert por["sinalizacao"]["tipo_fonte"] == "auxiliar"

def test_catalogo_render_e_md(tmp_path):
    from catalog import fontes
    cat = fontes.catalogo(medir=True)
    txt = fontes.render_texto(cat)
    assert "CATÁLOGO DE DADOS NACIONAIS" in txt and "IBGE" in txt
    out = tmp_path / "cat.md"
    got = fontes.gerar_md(out, cat)
    assert got == str(out) and out.exists()
    c = out.read_text(encoding="utf-8")
    assert c.startswith("# Catálogo de Dados Nacionais") and "Rastreabilidade" in c
    assert "BC250" in c                       # procedência real na tabela

def test_catalogo_cli():
    from catalog import fontes
    assert fontes._cli([]) == 0


# ============================================================================
#  MAPA DE COBERTURA NACIONAL (§41) — matriz UF × camada; gaps identificáveis.
# ============================================================================
def test_mapa_cobertura_html_estrutura():
    from audit import coverage_map as cm
    h = cm.construir_html(rapido=True)
    assert "<!doctype html>" in h and "Mapa de Cobertura Nacional" in h
    assert "prefers-color-scheme" in h                 # tema claro/escuro
    for regiao in ("Norte", "Nordeste", "Centro-Oeste", "Sudeste", "Sul"):
        assert regiao in h
    assert "class='cell ok" in h and "class='cell bad" in h   # presença e lacuna visíveis
    assert h.count("class='uf'") == 28                 # 27 UFs + cabeçalho

def test_mapa_cobertura_gera_arquivo(tmp_path):
    from audit import coverage_map as cm
    out = tmp_path / "mapa.html"
    got = cm.gerar(out, rapido=True)
    assert got == str(out) and out.exists()
    assert out.read_text(encoding="utf-8").startswith("<!doctype html>")

def test_mapa_cobertura_todas_as_ufs_presentes_no_html():
    from audit import coverage_map as cm
    h = cm.construir_html(rapido=True)
    for _, ufs in cm.REGIOES:
        for uf in ufs:
            assert (">%s<" % uf) in h


# ============================================================================
#  INTEGRIDADE GEOMÉTRICA (§20/§43) — defeitos reais nas feições instaladas.
# ============================================================================
def test_integridade_dados_reais_limpos():
    from audit import integridade as ig
    aud = ig.auditar()
    por = {c["chave"]: c for c in aud["camadas"]}
    # as bases nacionais embarcadas devem estar geometricamente limpas (0 inválidas/fora do BR)
    assert por["rodovias"]["invalidas"] == 0 and por["rodovias"]["fora_brasil"] == 0
    assert por["drenagem"]["feicoes"] > 2_000_000 and por["drenagem"]["invalidas"] == 0
    assert aud["resumo"]["total_invalidas"] == 0 and aud["resumo"]["total_fora_brasil"] == 0
    assert aud["resumo"]["status"] == "OK"

def test_integridade_detecta_defeitos_sinteticos(tmp_path, monkeypatch):
    # injeta uma camada com bbox invertida, coord nula e ponto fora do Brasil → deve acusar
    import numpy as np
    from audit import integridade as ig

    class _FakeRepo:
        def arrays(self, chave):
            return {"n": 4,
                    "lon": np.array([-47.0, np.nan, 10.0, -46.0]),    # 2ª nula, 3ª fora do BR
                    "lat": np.array([-15.0, -15.0, 10.0, -15.0]),
                    "xmin": np.array([-47.0, -47.0, 10.0, -46.1]),
                    "ymin": np.array([-15.1, -15.1, 9.9, -15.1]),
                    "xmax": np.array([-46.9, -46.9, 10.1, -47.0]),    # 4ª invertida (xmin>xmax)
                    "ymax": np.array([-14.9, -14.9, 10.1, -14.9])}
    r = ig.checar_camada(_FakeRepo(), "x")
    assert r["coord_nula"] == 1 and r["fora_brasil"] == 1 and r["invalidas"] >= 1
    assert r["status"] == "PARCIAL"

def test_integridade_camada_ausente(tmp_path):
    import local_data
    from audit import integridade as ig
    from geo import repositorio
    reg = local_data.LocalDataRegistry(cfg.app_root(), tmp_path / "dl")
    repo = repositorio.GeoIntelligenceRepository(registry=reg)
    r = ig.checar_camada(repo, "osrm_brasil")
    assert r["instalado"] is False and r["status"] == "AUSENTE"

def test_integridade_render_e_cli():
    from audit import integridade as ig
    aud = ig.auditar()
    txt = ig.render_texto(aud)
    assert "INTEGRIDADE GEOMÉTRICA" in txt and "rodovias" in txt
    assert ig._cli() == 0


# ============================================================================
#  DOSSIÊ DE ROTA — integração TOTAL dos dados na rota (§27/§28/§29/§30/§76).
# ============================================================================
@pytest.fixture(scope="module")
def dossie_sp_rio():
    from geo import dossie_rota
    # SP → Rio: rota real que cruza MG/RJ/SP, rios, pontes, balsa, ferrovias
    return dossie_rota.dossie([(-46.63, -23.55), (-43.20, -22.90)], folga_km=3.0)

def test_dossie_origem_destino_municipio(dossie_sp_rio):
    d = dossie_sp_rio
    assert d["origem"]["municipio"]["uf"] == "SP"
    assert d["destino"]["municipio"]["uf"] == "RJ"
    assert d["origem"]["municipio"]["nome"] and d["destino"]["municipio"]["nome"]

def test_dossie_travessia_territorial(dossie_sp_rio):
    tt = dossie_sp_rio["travessia_territorial"]
    assert tt["n_municipios"] > 0
    # o corredor SP→Rio passa por SP, MG e RJ
    assert {"SP", "RJ"}.issubset(set(tt["ufs"]))

def test_dossie_hidrografia_rica(dossie_sp_rio):
    h = dossie_sp_rio["hidrografia"]
    assert h["trechos"] > 100
    assert h["rios_nomeados"]                         # rios nomeados reais
    nomes = " ".join(r["nome"] for r in h["rios_nomeados"]).lower()
    assert "paraíba do sul" in nomes or "tietê" in nomes or "rio" in nomes

def test_dossie_camadas_e_alertas(dossie_sp_rio):
    cam = dossie_sp_rio["camadas"]
    assert cam["rodovias"]["feicoes"] > 0 and cam["pontes"]["feicoes"] >= 0
    assert isinstance(dossie_sp_rio["alertas"], list) and len(dossie_sp_rio["alertas"]) >= 1
    assert "drenagem" in dossie_sp_rio["fontes"]      # rastreabilidade das fontes

def test_dossie_multimodal_e_render(dossie_sp_rio):
    assert dossie_sp_rio["multimodal"].get("rodoviario", 0) > 0
    from geo import dossie_rota
    txt = dossie_rota.render_texto(dossie_sp_rio)
    assert "DOSSIÊ DA ROTA" in txt and "Hidrografia" in txt and "Alertas" in txt

def test_dossie_rota_vazia_e_parse():
    from geo import dossie_rota
    assert dossie_rota.dossie([]).get("erro")
    assert dossie_rota._parse_coords("-46.6,-23.5;-43.2,-22.9") == [(-46.6, -23.5), (-43.2, -22.9)]
    assert dossie_rota._cli([]) == 2                  # sem coords → uso

def test_dossie_ponto_unico_nao_quebra():
    from geo import dossie_rota
    from audit import coverage_auditor as ca
    d = dossie_rota.dossie([(-47.88, -15.79)])        # região de Brasília, ponto único
    # município mais próximo (nearest-centroid honesto): DF ou GO vizinho — UF válida
    assert d["origem"]["municipio"]["uf"] in ca.UFS_ESPERADAS and d["pontos"] == 1
    assert d["origem"]["municipio"]["nome"]


# ============================================================================
#  GRAFO TOPOLÓGICO / CONECTIVIDADE (§43) — nós/arestas via WKB + Union-Find.
# ============================================================================
def test_grafo_extremidades_wkb_linestring():
    import struct
    from geo import grafo_topologico as gt
    # LineString big-endian com 3 pontos: (1,2)->(3,4)->(5,6)
    pts = [(1.0, 2.0), (3.0, 4.0), (5.0, 6.0)]
    b = b"\x00" + struct.pack(">I", 2) + struct.pack(">I", 3)
    for x, y in pts:
        b += struct.pack(">dd", x, y)
    r = gt._extremidades_wkb(b)
    assert r == (1.0, 2.0, 5.0, 6.0)                   # início e fim

def test_grafo_extremidades_wkb_lixo_nao_quebra():
    from geo import grafo_topologico as gt
    assert gt._extremidades_wkb(b"\x00\x00") is None

def test_grafo_unionfind_componentes():
    from geo import grafo_topologico as gt
    uf = gt._UF(5)
    uf.union(0, 1); uf.union(1, 2); uf.union(3, 4)
    raizes = {uf.find(i) for i in range(5)}
    assert len(raizes) == 2                            # {0,1,2} e {3,4}

@pytest.fixture(scope="module")
def grafo_rodovias():
    from geo import grafo_topologico as gt
    return gt.construir("rodovias", precisao=3)

def test_grafo_rodovias_bem_conectado(grafo_rodovias):
    g = grafo_rodovias
    assert g["status"] == "OK"
    assert g["arestas"] > 250000 and g["nos"] > 100000
    # malha rodoviária nacional real: maior componente domina (muito acima de 90%)
    assert g["pct_maior"] > 90.0
    assert g["componentes"] >= 1 and g["nos_grau1"] >= 0

def test_grafo_render_e_cli(grafo_rodovias):
    from geo import grafo_topologico as gt
    txt = gt.render_texto(grafo_rodovias)
    assert "GRAFO TOPOLÓGICO" in txt and "maior componente" in txt
    assert gt._cli(["ferrovias"]) == 0                 # camada menor, roda rápido

def test_grafo_camada_ausente():
    from geo import grafo_topologico as gt
    g = gt.construir("osrm_brasil")
    assert g["status"] == "AUSENTE"


# ============================================================================
#  TELEMETRIA INCREMENTAL (§9/§10/§11/§34/§35) — store append-only + watermark +
#  dedup; laço incremental com fonte injetável (mock). Sem rede nos testes.
# ============================================================================
def test_telemetria_store_ingere_e_dedup(tmp_path):
    from telemetry import ana_incremental as tel
    st = tel.TelemetryStore(tmp_path / "tel")
    r1 = st.ingerir("estacoes", [{"codigo": "1", "nome": "A"}, {"codigo": "2", "nome": "B"}])
    assert r1 == {"novos": 2, "duplicados": 0, "total": 2}
    # reingerir 1 duplicado + 1 novo
    r2 = st.ingerir("estacoes", [{"codigo": "2", "nome": "B"}, {"codigo": "3", "nome": "C"}])
    assert r2["novos"] == 1 and r2["duplicados"] == 1 and r2["total"] == 3

def test_telemetria_watermark_persiste(tmp_path):
    from telemetry import ana_incremental as tel
    st = tel.TelemetryStore(tmp_path / "tel")
    assert st.carregar_watermark("chuva") is None
    st.salvar_watermark("chuva", "2026-10-05")
    assert tel.TelemetryStore(tmp_path / "tel").carregar_watermark("chuva") == "2026-10-05"

def test_telemetria_atualizar_incremental_so_novos(tmp_path):
    from telemetry import ana_incremental as tel
    st = tel.TelemetryStore(tmp_path / "tel")
    lote1 = [{"codigo": "9", "data": "2026-10-01", "cota": 100},
             {"codigo": "9", "data": "2026-10-02", "cota": 110}]
    chamadas = {"wm": []}
    def fetch(wm):
        chamadas["wm"].append(wm)
        # fonte "incremental": devolve tudo >= wm (simplificado devolve lote conforme wm)
        if wm is None:
            return lote1
        return [{"codigo": "9", "data": "2026-10-03", "cota": 120}]  # só o novo
    r1 = tel.atualizar_incremental("serie", fetch, store=st, campo_ts="data")
    assert r1["novos"] == 2 and r1["watermark_anterior"] is None and r1["watermark_novo"] == "2026-10-02"
    r2 = tel.atualizar_incremental("serie", fetch, store=st, campo_ts="data")
    assert r2["novos"] == 1 and r2["watermark_novo"] == "2026-10-03"   # avançou o watermark
    assert chamadas["wm"] == [None, "2026-10-02"]                      # 2ª busca partiu do watermark
    assert st.contar("serie") == 3

def test_telemetria_watermark_numerico(tmp_path):
    # Correção: timestamps NUMÉRICOS devem avançar por valor, não lexicograficamente ('9' vs '10').
    from telemetry import ana_incremental as tel
    st = tel.TelemetryStore(tmp_path / "telnum")
    lote = [{"id": "a", "data": "9"}, {"id": "b", "data": "10"}, {"id": "c", "data": "100"}]
    r = tel.atualizar_incremental("num", lambda wm: lote, store=st, campo_ts="data")
    assert r["watermark_novo"] == "100"            # 100 é o maior por VALOR (não '9' por texto)
    # e um batch com valor numérico maior avança o watermark (não trava)
    r2 = tel.atualizar_incremental("num", lambda wm: [{"id": "d", "data": "250"}],
                                   store=st, campo_ts="data")
    assert r2["watermark_novo"] == "250"
    assert tel._ordem_ts("100") > tel._ordem_ts("9")       # numérico por valor
    assert tel._ordem_ts("2026-10-03") > tel._ordem_ts("2026-10-02")   # ISO lexicográfico

def test_telemetria_total_sem_reler(tmp_path, monkeypatch):
    # Correção O(n²): ingerir não deve reler o arquivo inteiro (contar) para o total.
    from telemetry import ana_incremental as tel
    st = tel.TelemetryStore(tmp_path / "tel2")
    st.ingerir("s", [{"codigo": "1"}, {"codigo": "2"}])
    def _boom(*a, **k):
        raise AssertionError("contar() não deveria ser chamado em ingerir()")
    monkeypatch.setattr(st, "contar", _boom)
    r = st.ingerir("s", [{"codigo": "2"}, {"codigo": "3"}])   # 1 dup, 1 novo
    assert r["novos"] == 1 and r["duplicados"] == 1 and r["total"] == 3

def test_benchmark_idx_p95():
    from engines import benchmark as bm
    assert bm._idx_p95(2) == 1          # n=2 → o MAIOR (não o mínimo)
    assert bm._idx_p95(1) == 0
    assert bm._idx_p95(20) == 18        # ceil(0.95*20)-1
    assert bm._idx_p95(100) == 94
    assert bm._idx_p95(0) == 0

def test_repo_render_inventario_traco_sem_classes():
    # Correção de precedência: sem classes presentes, mostra '—' (não valor em branco).
    from geo import repositorio
    txt = repositorio.render_inventario({"camadas": {}, "classes_presentes": [], "total_feicoes": 0})
    assert "Classes multimodais presentes: —" in txt

def test_downloader_part_completo_recupera_em_416(tmp_path):
    # Correção: .part completo (processo morreu antes do rename) → 416 na retomada deve FINALIZAR
    # o arquivo, não falhar após N tentativas.
    from resources import downloader as dl
    import urllib.error
    destino = tmp_path / "arq.bin"
    part = tmp_path / "arq.bin.part"
    part.write_bytes(b"CONTEUDO-COMPLETO")        # .part já inteiro, aguardando rename
    def _abrir(url, inicio, timeout):
        raise urllib.error.HTTPError(url, 416, "Range Not Satisfiable", None, None)
    r = dl.baixar("http://x/arq.bin", destino, abrir=_abrir, tentativas=4,
                  dormir=lambda *_a: None)
    assert r["ok"] is True and destino.exists()
    assert destino.read_bytes() == b"CONTEUDO-COMPLETO" and not part.exists()

def test_telemetria_inventario(tmp_path):
    from telemetry import ana_incremental as tel
    st = tel.TelemetryStore(tmp_path / "tel")
    st.ingerir("estacoes", [{"codigo": "1"}])
    st.salvar_watermark("estacoes", "x")
    inv = st.inventario()
    assert inv["estacoes"]["registros"] == 1 and inv["estacoes"]["watermark"] == "x"

def test_telemetria_ana_client_sem_url_degrada():
    from telemetry import ana_incremental as tel
    cli = tel.AnaClient(conf={})            # sem telemetria.inventario_url
    r = cli.inventario_estacoes()
    assert r["ok"] is False and "inventario_url" in r["detalhe"]

def test_telemetria_ana_client_parse_csv_e_json():
    from telemetry import ana_incremental as tel
    cli = tel.AnaClient(conf={})
    assert cli._parse('[{"codigo":"1"},{"codigo":"2"}]') == [{"codigo": "1"}, {"codigo": "2"}]
    assert cli._parse('{"estacoes":[{"a":1}]}') == [{"a": 1}]
    linhas = cli._parse("codigo,nome\n1,A\n2,B")
    assert len(linhas) == 2 and linhas[0]["nome"] == "A"

def test_telemetria_ana_client_offline_no_net(monkeypatch):
    from telemetry import ana_incremental as tel
    monkeypatch.setenv("OPENROTAS_NO_NET", "1")
    cli = tel.AnaClient(conf={"telemetria": {"inventario_url": "http://x/export.json"}})
    r = cli.inventario_estacoes()
    assert r["ok"] is False and "rede desabilitada" in r["detalhe"]

def test_dossie_estacao_proxima_do_store(tmp_path, monkeypatch):
    # com inventário local, o dossiê surfaça a estação mais próxima (integração telemetria→rota)
    from telemetry import ana_incremental as tel
    from geo import dossie_rota
    import desktop_config as _cfg
    monkeypatch.setattr(_cfg, "user_data_dir", lambda: tmp_path / "ud")
    st = tel.TelemetryStore((tmp_path / "ud") / "telemetry")
    st.ingerir("estacoes", [{"nome": "Est. Tietê", "lat": -23.5, "lon": -46.6},
                            {"nome": "Est. Longe", "lat": 2.0, "lon": -60.0}])
    r = dossie_rota._estacao_proxima(-46.63, -23.55)
    assert r.get("estacao") == "Est. Tietê" and r.get("dist_km") is not None

def test_dossie_estacao_no_equador_lat_zero(tmp_path, monkeypatch):
    # Regressão: lat=0,0 (equador — cruza AP/PA/AM/RR) é falsy; `r.get("lat") or ...` a
    # descartaria e a estação seria pulada. O dossiê deve resolver a estação no equador.
    from telemetry import ana_incremental as tel
    from geo import dossie_rota
    import desktop_config as _cfg
    monkeypatch.setattr(_cfg, "user_data_dir", lambda: tmp_path / "ud")
    st = tel.TelemetryStore((tmp_path / "ud") / "telemetry")
    st.ingerir("estacoes", [{"nome": "Est. Equador", "lat": 0.0, "lon": -50.0}])
    r = dossie_rota._estacao_proxima(-50.0, 0.0)
    assert r.get("estacao") == "Est. Equador" and r.get("dist_km") is not None


# ============================================================================
#  DOSSIÊ: exportação HTML/Excel + surfacing no painel (§11/§12 — camada desktop).
# ============================================================================
def test_dossie_gera_html(tmp_path, dossie_sp_rio):
    from geo import dossie_rota
    out = tmp_path / "dossie.html"
    got = dossie_rota.gerar_html(dossie_sp_rio, out)
    assert got == str(out) and out.exists()
    c = out.read_text(encoding="utf-8")
    assert c.startswith("<!doctype html>") and "Dossiê da Rota" in c
    assert "prefers-color-scheme" in c and "Hidrografia" in c and "Alertas" in c

def test_dossie_exporta_excel(tmp_path, dossie_sp_rio):
    import openpyxl
    from geo import dossie_rota
    out = tmp_path / "dossie.xlsx"
    got = dossie_rota.exportar_excel(dossie_sp_rio, out)
    assert got == str(out) and out.exists()
    wb = openpyxl.load_workbook(out)
    assert set(["Resumo", "Rios", "Feicoes", "Fontes", "Alertas"]).issubset(set(wb.sheetnames))

def test_dossie_export_erro_nao_quebra(tmp_path):
    from geo import dossie_rota
    assert dossie_rota.gerar_html({"erro": "x"}, tmp_path / "a.html") is None
    assert dossie_rota.exportar_excel({"erro": "x"}, tmp_path / "a.xlsx") is None

def test_painel_tem_bloco_dossie():
    h = painel.construir_html()
    assert "Dossiê de Rota (integração total)" in h and "--dossie" in h
    assert "Comparação de Rotas" in h and "--comparar" in h   # nova capacidade surfaçada no painel


# ============================================================================
#  COMPARAÇÃO DE ROTAS — ranqueia rotas candidatas por complexidade logística (desktop).
# ============================================================================
def _dd_cmp(travessias=0, pontes=0, trechos=0, naveg=0, ufs=("SP",), o="A", d="B", erro=False, naveg_nome=0):
    """Dossiê SINTÉTICO (sem tocar nas bases) para testar a pontuação/ranking de forma determinística."""
    if erro:
        return {"erro": "x"}
    return {
        "origem": {"municipio": {"nome": o, "uf": (ufs[0] if ufs else "?")}},
        "destino": {"municipio": {"nome": d, "uf": (ufs[-1] if ufs else "?")}},
        "camadas": {"travessias": {"feicoes": travessias}, "pontes": {"feicoes": pontes}},
        "hidrografia": {"trechos": trechos, "navegaveis": naveg, "navegaveis_nomeados": naveg_nome},
        "travessia_territorial": {"ufs": list(ufs)},
        "pontos": 2,
    }

def test_comparar_pontuar_formula_e_erro():
    from geo import dossie_comparativo as dc
    s, c = dc.pontuar(_dd_cmp(travessias=1, pontes=2, trechos=100, naveg=3, ufs=("SP", "MG", "RJ")))
    assert c == {"travessias": 1, "transicoes_uf": 2, "rios_navegaveis": 3, "pontes": 2, "trechos_rio": 100}
    # 10*1 + 3*2 + 2*3 + 1*2 + 0.05*100 = 29.0
    assert s == round(10 * 1 + 3 * 2 + 2 * 3 + 1 * 2 + 0.05 * 100, 2)
    assert dc.pontuar({"erro": "x"}) == (None, {})
    assert dc.pontuar({}) == (None, {})

def test_comparar_pontuar_usa_navegavel_por_nome():
    # A base pode trazer navegaveis=0 (desconhecido), mas um rio navegável reconhecido por NOME (ex.:
    # Amazonas) DEVE pesar na complexidade — senão a rota que cruza o Amazonas pontua como se fosse seca.
    from geo import dossie_comparativo as dc
    s0, c0 = dc.pontuar(_dd_cmp(naveg=0, naveg_nome=0))
    s1, c1 = dc.pontuar(_dd_cmp(naveg=0, naveg_nome=1))        # Amazonas reconhecido por nome
    assert c0["rios_navegaveis"] == 0 and c1["rios_navegaveis"] == 1
    assert s1 - s0 == 2.0                                      # peso 2,0 por rio navegável
    # a base continua valendo quando é MAIOR que o reconhecido por nome (não regride).
    _, c2 = dc.pontuar(_dd_cmp(naveg=3, naveg_nome=1))
    assert c2["rios_navegaveis"] == 3

def test_comparar_ranqueia_e_recomenda(monkeypatch):
    from geo import dossie_rota, dossie_comparativo as dc
    fila = [
        _dd_cmp(travessias=2, naveg=1, ufs=("SP", "MG", "RJ"), o="SP", d="RJ"),   # complexa
        _dd_cmp(travessias=0, naveg=0, ufs=("SP",), o="SP", d="Campinas"),        # simples
    ]
    estado = {"i": 0}
    def _fake(coords, **k):
        r = fila[estado["i"]]
        estado["i"] += 1
        return r
    monkeypatch.setattr(dossie_rota, "dossie", _fake)
    cmp = dc.comparar([[(0, 0), (1, 1)], [(0, 0), (0.1, 0.1)]], nomes=["Rota longa", "Rota curta"])
    assert cmp["total"] == 2 and cmp["validas"] == 2
    # a rota mais simples (menor score) vai para o 1º lugar e é a recomendada
    assert cmp["rotas"][0]["componentes"]["travessias"] == 0
    assert cmp["rotas"][0]["score"] < cmp["rotas"][1]["score"]
    assert "Campinas" in cmp["recomendada"]
    txt = dc.render_texto(cmp)
    assert "COMPARAÇÃO DE ROTAS" in txt and "RECOMENDADA" in txt

def test_comparar_rota_invalida_sai_do_ranking(monkeypatch):
    from geo import dossie_rota, dossie_comparativo as dc
    fila = [_dd_cmp(erro=True), _dd_cmp(travessias=1, ufs=("BA",), o="X", d="Y")]
    estado = {"i": 0}
    monkeypatch.setattr(dossie_rota, "dossie", lambda c, **k: fila[estado.__setitem__("i", estado["i"] + 1) or estado["i"] - 1])
    cmp = dc.comparar([[(0, 0), (1, 1)], [(2, 2), (3, 3)]])
    assert cmp["total"] == 2 and cmp["validas"] == 1
    # a válida vem primeiro; a inválida fica por último e sem score
    assert cmp["rotas"][0]["ok"] is True and cmp["rotas"][-1]["ok"] is False
    assert cmp["rotas"][-1]["score"] is None

def test_comparar_html_e_excel(tmp_path, monkeypatch):
    import openpyxl
    from geo import dossie_rota, dossie_comparativo as dc
    fila = [_dd_cmp(travessias=0, ufs=("SP",), o="SP", d="Campinas"),
            _dd_cmp(travessias=2, ufs=("SP", "RJ"), o="SP", d="RJ")]
    estado = {"i": 0}
    monkeypatch.setattr(dossie_rota, "dossie", lambda c, **k: fila[estado.__setitem__("i", estado["i"] + 1) or estado["i"] - 1])
    cmp = dc.comparar([[(0, 0), (1, 1)], [(0, 0), (2, 2)]])
    h = dc.gerar_html(cmp, tmp_path / "cmp.html")
    assert h == str(tmp_path / "cmp.html") and (tmp_path / "cmp.html").exists()
    c = (tmp_path / "cmp.html").read_text(encoding="utf-8")
    assert c.startswith("<!doctype html>") and "Comparação de Rotas" in c
    assert "prefers-color-scheme" in c and "RECOMENDADA" in c
    x = dc.exportar_excel(cmp, tmp_path / "cmp.xlsx")
    assert x == str(tmp_path / "cmp.xlsx") and (tmp_path / "cmp.xlsx").exists()
    wb = openpyxl.load_workbook(tmp_path / "cmp.xlsx")
    assert "Comparacao" in wb.sheetnames

def test_prov_comparar_rotas_grava_exports(monkeypatch, tmp_path):
    import sys as _sys, pathlib as _pl
    _sys.path.insert(0, str(_pl.Path(__file__).resolve().parents[1]))
    from geo import dossie_rota
    monkeypatch.setattr(prov, "pasta_por_tipo", lambda tipo: tmp_path / tipo)
    fila = [_dd_cmp(travessias=0, ufs=("SP",)), _dd_cmp(travessias=3, ufs=("SP", "RJ"))]
    estado = {"i": 0}
    monkeypatch.setattr(dossie_rota, "dossie", lambda c, **k: fila[estado.__setitem__("i", estado["i"] + 1) or estado["i"] - 1])
    r = prov.comparar_rotas([[(0, 0), (1, 1)], [(0, 0), (2, 2)]])
    assert r["ok"] is True and r["validas"] == 2
    assert r["html"] and "exports" in r["html"] and r["html"].endswith(".html")
    assert r["excel"] and r["excel"].endswith(".xlsx")

def test_comparar_cli_coords(monkeypatch, capsys):
    from geo import dossie_rota, dossie_comparativo as dc
    fila = [_dd_cmp(travessias=0, ufs=("SP",), o="SP", d="Campinas"),
            _dd_cmp(travessias=2, ufs=("SP", "RJ"), o="SP", d="RJ")]
    estado = {"i": 0}
    monkeypatch.setattr(dossie_rota, "dossie", lambda c, **k: fila[estado.__setitem__("i", estado["i"] + 1) or estado["i"] - 1])
    rc = dc._cli(["-46.6,-23.5;-47.0,-22.9 || -46.6,-23.5;-43.2,-22.9"])
    assert rc == 0
    out = capsys.readouterr().out
    assert "COMPARAÇÃO DE ROTAS" in out
    # menos de 2 rotas → uso/erro
    assert dc._cli(["-46.6,-23.5;-47.0,-22.9"]) == 2


# ============================================================================
#  TEMA DA MARCA — sistema de design único dos relatórios + janela (visual coeso).
# ============================================================================
def test_tema_css_tem_tokens_e_dois_temas():
    from resources import tema
    c = tema.css()
    # tokens da marca e os DOIS temas (3 estados) presentes
    assert "--accent" in c and "--teal" in c and "--surface" in c
    assert "prefers-color-scheme:dark" in c
    assert ':root:not([data-theme="light"])' in c and ':root[data-theme="dark"]' in c
    # offline-first: nada de fontes/recursos externos
    assert "fonts.googleapis" not in c and "http://" not in c and "https://" not in c

def test_tema_logo_e_documento():
    from resources import tema
    assert tema.logo_svg().startswith("<svg") and "</svg>" in tema.logo_svg()
    doc = tema.documento("Título X", "sub", "<p>corpo</p>", rodape="rod")
    assert doc.startswith("<!doctype html>") and "Título X" in doc and tema.MARCA in doc
    assert "corpo" in doc and "prefers-color-scheme" in doc

def test_tema_tk_defensivo_nao_levanta():
    from resources import tema
    # root inválido → cai no except e devolve False, sem levantar (contrato "nunca quebra")
    assert tema.aplicar_tema_tk(object()) is False

def test_relatorios_usam_tema_compartilhado(tmp_path, monkeypatch, dossie_sp_rio):
    # Todos os relatórios devem embutir o tema da marca (token --teal só existe no tema novo,
    # não nos CSS de fallback antigos) — prova visual de que a identidade única está aplicada.
    from resources import painel
    from audit import saude_nacional as sn
    from audit import coverage_map as cm
    from geo import dossie_rota, dossie_comparativo as dc
    assert "--teal" in painel.construir_html()
    assert "--teal" in sn.construir_html(rapido=True)
    assert "--teal" in cm.construir_html(rapido=True)
    out = tmp_path / "d.html"
    dossie_rota.gerar_html(dossie_sp_rio, out)
    assert "--teal" in out.read_text(encoding="utf-8")
    from geo import dossie_rota as _dr
    fila = [_dd_cmp(travessias=0, ufs=("SP",)), _dd_cmp(travessias=2, ufs=("SP", "RJ"))]
    estado = {"i": 0}
    monkeypatch.setattr(_dr, "dossie", lambda c, **k: fila[estado.__setitem__("i", estado["i"] + 1) or estado["i"] - 1])
    cmp = dc.comparar([[(0, 0), (1, 1)], [(0, 0), (2, 2)]])
    dc.gerar_html(cmp, tmp_path / "c.html")
    assert "--teal" in (tmp_path / "c.html").read_text(encoding="utf-8")

def test_tema_grafico_barras_e_proporcao():
    from resources import tema
    g = tema.grafico_barras([{"rotulo": "A", "valor": 3},
                             {"rotulo": "B", "valor": 9, "classe": "ok"}], unidade="ms", titulo="T")
    assert "<svg class='ork-chart barras'" in g and "data-ork-tip" in g
    assert "class='bar ok'" in g and "ork-tip" in g      # série colorida por estado + tooltip JS
    assert tema.grafico_barras([]) == "" and tema.grafico_barras(None) == ""
    p = tema.barra_proporcao([{"rotulo": "OK", "valor": 5, "classe": "ok"},
                              {"rotulo": "Ruim", "valor": 2, "classe": "bad"}], titulo="P")
    assert "ork-chart" in p and "ork-leg" in p and "data-ork-tip" in p
    assert tema.barra_proporcao([]) == ""

def test_comparar_html_tem_grafico(tmp_path, monkeypatch):
    from geo import dossie_rota, dossie_comparativo as dc
    fila = [_dd_cmp(travessias=0, ufs=("SP",), o="SP", d="Campinas"),
            _dd_cmp(travessias=2, ufs=("SP", "RJ"), o="SP", d="RJ")]
    estado = {"i": 0}
    monkeypatch.setattr(dossie_rota, "dossie", lambda c, **k: fila[estado.__setitem__("i", estado["i"] + 1) or estado["i"] - 1])
    cmp = dc.comparar([[(0, 0), (1, 1)], [(0, 0), (2, 2)]])
    dc.gerar_html(cmp, tmp_path / "c.html")
    h = (tmp_path / "c.html").read_text(encoding="utf-8")
    assert "ork-chart" in h and "data-ork-tip" in h and "Ranking de complexidade" in h

def test_dossie_html_tem_grafico(tmp_path, dossie_sp_rio):
    from geo import dossie_rota
    dossie_rota.gerar_html(dossie_sp_rio, tmp_path / "d.html")
    h = (tmp_path / "d.html").read_text(encoding="utf-8")
    assert "ork-chart" in h and "FEIÇÕES POR CAMADA" in h

def test_saude_html_tem_grafico_proporcao():
    from audit import saude_nacional as sn
    h = sn.construir_html(rapido=True)
    assert "ork-chart" in h and "ork-leg" in h and "DIMENSÕES POR ESTADO" in h


# ============================================================================
#  MAPA NACIONAL (dispersão de municípios, SVG) — relatórios + ícone da marca.
# ============================================================================
def test_tema_mapa_dispersao():
    from resources import tema
    xs = [-46, -43, -60, -38]
    ys = [-23, -22, -3, -8]
    m = tema.mapa_dispersao(xs, ys, classes=["teal", None, "bad", None],
                            rota=[(-46, -23), (-43, -22)], titulo="MAPA",
                            legenda=[{"rotulo": "x", "cor": "var(--teal)"}])
    assert "<svg class='ork-mapa'" in m and "class='pt teal'" in m and "class='pt bad'" in m
    assert "class='rota'" in m and "class='wp'" in m and "ork-leg" in m
    assert tema.mapa_dispersao([], []) == "" and tema.mapa_dispersao(None, None) == ""

def test_coverage_pontos_municipios():
    from audit import coverage_auditor as ca
    lons, lats, ufs = ca.pontos_municipios()
    assert len(lons) == len(lats) == len(ufs) and len(lons) > 0
    assert any(u in ("SP", "RJ", "AM", "BA") for u in ufs if u)   # UFs reais presentes

def test_coverage_map_tem_mapa_nacional():
    from audit import coverage_map as cm
    h = cm.construir_html(rapido=True)
    assert "ork-mapa" in h and "MAPA NACIONAL" in h

def test_dossie_html_tem_mapa_nacional(tmp_path, dossie_sp_rio):
    from geo import dossie_rota
    dossie_rota.gerar_html(dossie_sp_rio, tmp_path / "dm.html")
    h = (tmp_path / "dm.html").read_text(encoding="utf-8")
    assert "ork-mapa" in h and "LOCALIZAÇÃO DA ROTA" in h and "class='rota'" in h

def test_make_icon_render_marca_de_rota():
    import make_icon
    try:
        from PIL import Image  # noqa: F401
    except Exception:
        import pytest
        pytest.skip("Pillow ausente")
    img = make_icon._render(64)
    assert img.size == (64, 64) and img.mode == "RGBA"


# ============================================================================
#  POLIMENTO: alternador de tema, animações de entrada e mapa coroplético por UF.
# ============================================================================
def test_tema_ui_js_toggle():
    from resources import tema
    j = tema.ui_js()
    assert "<script>" in j and "ork-theme-toggle" in j and "localStorage" in j and "data-theme" in j

def test_tema_css_toggle_e_animacoes():
    from resources import tema
    c = tema.css()
    assert ".ork-theme-toggle" in c
    assert "@keyframes ork-grow" in c and "@keyframes ork-draw" in c and "@keyframes ork-fade" in c
    assert "prefers-reduced-motion:no-preference" in c          # animações só quando permitido

def test_tema_mapa_coropletico():
    from resources import tema
    polis = [{"uf": "SP", "pts": [(-53, -25), (-44, -20), (-48, -19), (-50, -24)], "classe": "teal", "dica": "SP"},
             {"uf": "RJ", "pts": [(-44, -23), (-41, -21), (-43, -22)], "classe": "bad"}]
    m = tema.mapa_coropletico(polis, titulo="UF", rota=[(-46, -23), (-43, -22)],
                              legenda=[{"rotulo": "x", "cor": "var(--teal)"}])
    assert "<svg class='ork-mapa'" in m and "class='uf teal'" in m and "uf-lbl" in m
    assert "class='rota'" in m and "ork-leg" in m and "ork-theme-toggle" in m
    assert tema.mapa_coropletico([]) == "" and tema.mapa_coropletico(None) == ""

def test_coverage_hulls_uf():
    from audit import coverage_auditor as ca
    h = ca.hulls_uf()
    assert isinstance(h, dict) and len(h) >= 20            # maioria das 27 UFs
    assert "SP" in h and all(len(v) >= 3 for v in h.values())

def test_coverage_map_choropleth_por_uf():
    from audit import coverage_map as cm
    h = cm.construir_html(rapido=True)
    assert "MAPA NACIONAL POR UF" in h and "class='uf " in h and "ork-mapa" in h

def test_relatorios_tem_alternador_de_tema():
    from audit import saude_nacional as sn
    from resources import painel
    assert "ork-theme-toggle" in sn.construir_html(rapido=True)
    assert "ork-theme-toggle" in painel.construir_html()

def test_tema_favicon_e_print():
    from resources import tema
    j = tema.ui_js()
    assert "rel='icon'" in j or "rel=\\\"icon\\\"" in j or "link" in j
    assert "image/svg" in j                         # favicon da marca (data-URI)
    assert "@media print" in tema.css()             # estilo de impressão/PDF

def test_painel_executivo_html():
    from resources import painel_executivo as pe
    from resources import tema
    h = pe.construir_html()
    assert h.startswith("<!doctype html>") and "Painel Executivo" in h and tema.MARCA in h
    assert "Estados (UFs)" in h and "Feições multimodais" in h
    assert "COBERTURA NACIONAL POR UF" in h          # mapa coroplético embutido
    assert "ork-theme-toggle" in h and "prefers-color-scheme" in h

def test_painel_executivo_gera(tmp_path):
    from resources import painel_executivo as pe
    out = tmp_path / "pe.html"
    assert pe.gerar(out) == str(out) and out.exists()
    assert out.read_text(encoding="utf-8").startswith("<!doctype html>")


# ============================================================================
#  CONECTIVIDADE MULTIMODAL REAL (§8) — nós compartilhados entre modais (por célula).
# ============================================================================
def test_multimodal_analisar():
    import itertools
    from geo import multimodal
    g = multimodal.analisar(precisao=1)
    assert g["status"] == "OK" and "rodoviario" in g["classes"]
    assert g["total_celulas"] > 0 and g["hubs"] >= 0 and g["hubs3"] >= 0
    # diagonal (células próprias) >= qualquer co-ocorrência da linha; matriz simétrica
    for a in g["classes"]:
        assert g["matriz"][a][a] >= max(g["matriz"][a][b] for b in g["classes"])
    for a, b in itertools.combinations(g["classes"], 2):
        assert g["matriz"][a][b] == g["matriz"][b][a]
    # transposições, quando presentes, têm pct coerente (0..100)
    for t in g["transposicoes"].values():
        assert 0 <= t["conecta"] <= t["total"] and 0.0 <= t["pct"] <= 100.0
    assert "CONECTIVIDADE MULTIMODAL" in multimodal.render_texto(g)

def test_multimodal_html_e_gera(tmp_path):
    from geo import multimodal
    h = multimodal.construir_html()
    assert h.startswith("<!doctype html>") and "Conectividade Multimodal" in h
    assert "Hubs multimodais" in h and "ork-theme-toggle" in h
    out = tmp_path / "mm.html"
    assert multimodal.gerar(out) == str(out) and out.exists()


# ============================================================================
#  ISÓCRONAS DE ALCANCE (§29) — municípios atingíveis por faixa de distância.
# ============================================================================
def test_isocronas_calcular_estimado():
    # [PARIDADE-WEB] Sem OSRM, o padrão agora ESTIMA a distância rodoviária a partir da linha reta
    # (fator de sinuosidade do web), em vez de reportar a reta crua — a aplicação web nunca reporta a
    # reta como se fosse a viagem. Modo 'estimado'; as mesmas invariantes de faixa/cumulativo se mantêm.
    from geo import isocronas
    r = isocronas.calcular((-46.63, -23.55), faixas_km=(100, 300, 600, 1000))
    assert r["ok"] and r["modo"] == "estimado"
    assert len(r["bandas"]) == 4 and r["total_municipios"] > 0
    cums = [b["cumulativo"] for b in r["bandas"]]
    assert cums == sorted(cums)                                   # cumulativo monotônico
    assert sum(b["municipios"] for b in r["bandas"]) == r["alcancados"] == cums[-1]
    assert r["alcancados"] <= r["total_municipios"]
    assert r["bandas"][0]["municipios"] > 0                       # vizinhos imediatos de SP
    assert "ISÓCRONAS" in isocronas.render_texto(r)
    # estimar=False mantém a LINHA RETA pura (piso físico): modo 'geodesico' e alcance MAIOR (reta < viária).
    rg = isocronas.calcular((-46.63, -23.55), faixas_km=(100, 300, 600, 1000), estimar=False)
    assert rg["ok"] and rg["modo"] == "geodesico"
    assert rg["alcancados"] >= r["alcancados"]                    # a estimativa empurra municípios p/ além

def test_isocronas_osrm_cai_em_estimado_sem_rede(monkeypatch):
    # Sem rede, o OSRM é pulado e o padrão cai na ESTIMATIVA rodoviária (não mais na reta crua).
    monkeypatch.setenv("OPENROTAS_NO_NET", "1")
    from geo import isocronas
    r = isocronas.calcular((-46.63, -23.55), usar_osrm=True)
    assert r["ok"] and r["modo"] == "estimado"                    # OSRM pulado → estimativa de sinuosidade
    rg = isocronas.calcular((-46.63, -23.55), usar_osrm=True, estimar=False)
    assert rg["ok"] and rg["modo"] == "geodesico"                 # escape honesto p/ a linha reta pura

def test_isocronas_origem_por_nome():
    from geo import isocronas
    r = isocronas.calcular("São Paulo/SP", faixas_km=(100, 300))
    assert r["ok"] and r["origem"]["uf"] == "SP" and len(r["bandas"]) == 2

def test_isocronas_html_e_gera(tmp_path):
    from geo import isocronas
    r = isocronas.calcular((-47.88, -15.79))                      # Brasília
    h = isocronas.construir_html(r)
    assert h.startswith("<!doctype html>") and "Isócronas de Alcance" in h
    assert "ork-mapa" in h and "ork-theme-toggle" in h and "MUNICÍPIOS POR FAIXA" in h
    out = tmp_path / "iso.html"
    assert isocronas.gerar(out, r) == str(out) and out.exists()

def test_isocronas_origem_ambigua_avisa():
    # Origem homônima sem UF ('Bom Jesus') → calcula a partir da 1ª, mas avisa (texto e HTML).
    from geo import isocronas
    r = isocronas.calcular("Bom Jesus", faixas_km=(100, 300))
    assert r["ok"] and r["origem"].get("ambiguo") and len(r["origem"].get("ufs_alt", [])) >= 2
    assert "⚠" in isocronas.render_texto(r) and "ambígua" in isocronas.construir_html(r)
    r2 = isocronas.calcular("Bom Jesus/PI", faixas_km=(100, 300))
    assert not r2["origem"].get("ambiguo") and "⚠" not in isocronas.render_texto(r2)


# ============================================================================
#  MATRIZ ORIGEM↔DESTINO (§29) — distâncias entre vários centros.
# ============================================================================
def test_diagnostico_roteamento_offline():
    # [EVIDÊNCIA] Sem OSRM no ar: modo padrão 'estimado', amostra traz reta < estimada (fator ≥ 1) e real=None.
    from geo import dossie_rota as dr
    import os
    os.environ.pop("OSRM_URL", None)
    dr._OSRM_ATIVO_CACHE.update({"t": 0.0, "url": None})
    d = dr.diagnostico_roteamento([("São Paulo/SP", "Rio de Janeiro/RJ")])
    assert d["osrm_ativo"] is False and "estimado" in d["modo_padrao"]
    a = d["amostra"][0]
    assert a["de"] == "São Paulo/SP" and a["real_km"] is None
    assert a["reta_km"] and a["estimado_km"] and a["estimado_km"] >= a["reta_km"]   # estimativa ≥ reta
    txt = dr.render_diagnostico_roteamento(d)
    assert "DIAGNÓSTICO DE ROTEAMENTO" in txt and "não detectado" in txt

def test_diagnostico_roteamento_com_osrm(monkeypatch):
    # Com OSRM no ar (dublado), a amostra traz real_km e o ganho % vs a estimativa.
    from geo import dossie_rota as dr, matriz_od
    monkeypatch.setattr(dr, "osrm_ativo", lambda *a, **k: "http://localhost:5000")
    def _fake_table(lons, lats, url, **k):
        import numpy as np
        from geo import isocronas
        n = len(lons); M = [[0.0] * n for _ in range(n)]
        for i in range(n):
            row = isocronas._haversine_km(lons[i], lats[i], np.array(lons), np.array(lats))
            for j in range(n):
                M[i][j] = float(row[j]) * 1.2        # rodoviária real = 1,2× a reta (< estimativa ~1,25)
        return M
    monkeypatch.setattr(matriz_od, "_matriz_osrm", _fake_table)
    d = dr.diagnostico_roteamento([("São Paulo/SP", "Rio de Janeiro/RJ")])
    assert d["osrm_ativo"] is True and "rodoviário" in d["modo_padrao"]
    a = d["amostra"][0]
    assert a["real_km"] is not None and a["ganho_vs_estimativa_pct"] is not None

def test_osrm_auto_resolver_osrm():
    # Semântica do AUTO: None usa o OSRM local SE no ar; True força o configurado; False nunca.
    from geo import dossie_rota as dr
    import os
    # None sem OSRM no ar → (False, None); respeita OPENROTAS_NO_NET.
    assert dr.resolver_osrm(False) == (False, None)
    os.environ["OPENROTAS_NO_NET"] = "1"
    dr._OSRM_ATIVO_CACHE.update({"t": 0.0, "url": None})      # invalida cache
    try:
        assert dr.osrm_ativo() is None                        # sem rede → None
        assert dr.resolver_osrm(None) == (False, None)
    finally:
        del os.environ["OPENROTAS_NO_NET"]

def test_osrm_url_le_ambiente_e_modo_nativo(monkeypatch):
    # [ATIVAÇÃO AUTOMÁTICA] O launcher sobe o motor local (nativo 'auto'/'local') e injeta OSRM_URL no
    # ambiente. _osrm_url DEVE enxergar isso (era o elo que faltava) e reconhecer os modos que servem local.
    from geo import dossie_rota as dr
    monkeypatch.delenv("OSRM_URL", raising=False)
    # modo nativo 'auto'/'local' na config → reconhecido (default localhost:5000 / url configurada)
    assert dr._osrm_url({"osrm": {"mode": "auto"}}) == "http://localhost:5000"
    assert dr._osrm_url({"osrm": {"mode": "local", "url": "http://127.0.0.1:5000"}}) == "http://127.0.0.1:5000"
    assert dr._osrm_url({"osrm": {"mode": "off"}}) is None
    # OSRM_URL do ambiente tem PRIORIDADE (motor que o launcher subiu nesta sessão)
    monkeypatch.setenv("OSRM_URL", "http://localhost:5000/")
    assert dr._osrm_url({"osrm": {"mode": "off"}}) == "http://localhost:5000"

def test_matriz_auto_usa_osrm_quando_no_ar(monkeypatch):
    # [OSRM-AUTO] Com o OSRM local no ar, a matriz usa distância RODOVIÁRIA REAL por PADRÃO (sem flag) —
    # o ponto que o usuário cobrava ("já que tem roteamento OSRM local"). Dubla a disponibilidade + o /table.
    from geo import dossie_rota as dr, matriz_od
    monkeypatch.setattr(dr, "osrm_ativo", lambda *a, **k: "http://localhost:5000")
    # /table dublado: rodoviário ~1.35x a geodésica (determinístico, sem rede)
    def _fake_table(lons, lats, url, **k):
        from geo import isocronas
        import numpy as np
        n = len(lons)
        M = [[0.0] * n for _ in range(n)]
        for i in range(n):
            row = isocronas._haversine_km(lons[i], lats[i], np.array(lons), np.array(lats))
            for j in range(n):
                M[i][j] = float(row[j]) * 1.35
        return M
    monkeypatch.setattr(matriz_od, "_matriz_osrm", _fake_table)
    r = matriz_od.calcular([(-46.63, -23.55), (-43.20, -22.90)])   # SEM usar_osrm → AUTO
    assert r["ok"] and r["modo"] == "rodoviario"                   # usou o OSRM local automaticamente
    # usar_osrm=False continua forçando a estimativa (não toca no OSRM).
    rf = matriz_od.calcular([(-46.63, -23.55), (-43.20, -22.90)], usar_osrm=False)
    assert rf["modo"] == "estimado"

def test_matriz_osrm_celula_sem_rota_cai_na_estimativa(monkeypatch):
    # [OSRM-PARCIAL] Quando o /table devolve NULL para um par (ilha/fluvial sem estrada), a célula não pode
    # ficar com a LINHA RETA crua — cai na ESTIMATIVA (fluvial-aware). Modo vira 'rodoviario-parcial'.
    from geo import dossie_rota as dr, matriz_od
    monkeypatch.setattr(dr, "osrm_ativo", lambda *a, **k: "http://localhost:5000")
    def _fake_table(lons, lats, url, **k):
        # 2x2: par 0-1 roteado (real), par com índice 1 para 0 NULL (sem rota) — simula trecho sem estrada.
        return [[0.0, 120.0], [None, 0.0]]
    monkeypatch.setattr(matriz_od, "_matriz_osrm", _fake_table)
    r = matriz_od.calcular([(-46.63, -23.55), (-43.20, -22.90)])
    assert r["ok"] and r["modo"] == "rodoviario-parcial"       # teve rota real E célula estimada
    assert r["matriz"][0][1] == 120.0                          # célula roteada pelo OSRM
    # a célula NULL (1→0) não é a reta crua: recebeu a estimativa (reta × fator ≥ 1,18) > reta geodésica.
    import numpy as np
    from geo import isocronas
    reta10 = float(isocronas._haversine_km(-43.20, -22.90, np.array([-46.63]), np.array([-23.55]))[0])
    assert r["matriz"][1][0] > reta10 * 1.1

def test_matriz_calcular_estimado():
    # [PARIDADE-WEB] Sem OSRM, a matriz vira ESTIMATIVA rodoviária (linha reta × sinuosidade), como o web;
    # simetria e diâmetro continuam válidos. estimar=False recupera a linha reta pura (modo 'geodesico').
    from geo import matriz_od
    r = matriz_od.calcular([(-46.63, -23.55), (-43.20, -22.90), (-47.88, -15.79)])
    assert r["ok"] and r["modo"] == "estimado" and len(r["pontos"]) == 3
    M = r["matriz"]
    assert len(M) == 3 and all(len(row) == 3 for row in M)
    for i in range(3):
        assert M[i][i] == 0.0                              # diagonal zero preservada
        for j in range(3):
            assert abs(M[i][j] - M[j][i]) < 0.2            # simétrica
    rs = r["resumo"]
    assert rs["par_distante"]["km"] >= rs["par_proximo"]["km"] > 0
    assert rs["diametro_km"] == rs["par_distante"]["km"]
    assert "MATRIZ ORIGEM" in matriz_od.render_texto(r)
    # a estimativa é SEMPRE >= a linha reta (fator >= 1): o diâmetro estimado supera o geodésico.
    rg = matriz_od.calcular([(-46.63, -23.55), (-43.20, -22.90), (-47.88, -15.79)], estimar=False)
    assert rg["modo"] == "geodesico" and rs["diametro_km"] > rg["resumo"]["diametro_km"]

def test_matriz_menos_de_2_pontos():
    from geo import matriz_od
    assert matriz_od.calcular([(-46.63, -23.55)])["ok"] is False

def test_matriz_osrm_cai_em_estimado_sem_rede(monkeypatch):
    monkeypatch.setenv("OPENROTAS_NO_NET", "1")
    from geo import matriz_od
    r = matriz_od.calcular([(-46.63, -23.55), (-43.20, -22.90)], usar_osrm=True)
    assert r["ok"] and r["modo"] == "estimado"        # OSRM pulado sem rede → estimativa rodoviária

def test_matriz_html_e_excel(tmp_path):
    import openpyxl
    from geo import matriz_od
    r = matriz_od.calcular("São Paulo/SP;Rio de Janeiro/RJ;Curitiba/PR")
    assert r["ok"] and len(r["pontos"]) == 3
    h = matriz_od.construir_html(r)
    assert h.startswith("<!doctype html>") and "Matriz Origem" in h
    assert "ork-theme-toggle" in h and "Matriz de distâncias" in h
    out = tmp_path / "m.html"
    assert matriz_od.gerar(out, r) == str(out) and out.exists()
    xo = tmp_path / "m.xlsx"
    assert matriz_od.exportar_excel(r, xo) == str(xo) and xo.exists()
    assert "Matriz_km" in openpyxl.load_workbook(xo).sheetnames

def test_matriz_cidade_ambigua_avisa():
    # HONESTIDADE: 'Bom Jesus' sem UF é homônima em 5 UFs — a matriz resolve pela 1ª, mas o relatório
    # DEVE avisar (texto e HTML) para o usuário não computar a matriz da cidade errada em silêncio.
    from geo import matriz_od
    r = matriz_od.calcular("Bom Jesus;São Paulo/SP")
    assert r["ok"] and r["resumo"]["ambiguos"] == 1
    amb = [p for p in r["pontos"] if p.get("ambiguo")]
    assert len(amb) == 1 and len(amb[0]["ufs_alt"]) >= 2
    assert "⚠" in matriz_od.render_texto(r) and "ambígua" in matriz_od.construir_html(r)
    # com UF explícita não há ambiguidade nem aviso.
    r2 = matriz_od.calcular("Bom Jesus/PI;São Paulo/SP")
    assert r2["resumo"]["ambiguos"] == 0 and "⚠" not in matriz_od.render_texto(r2)

def test_matriz_local_nao_resolvido_avisa():
    # Honestidade: um local que não resolve é DROPADO da matriz — o relatório deve dizer, não omitir.
    from geo import matriz_od
    r = matriz_od.calcular("São Paulo/SP;CidadeInexistenteZZZ;Rio de Janeiro/RJ")
    assert r["ok"] and len(r["pontos"]) == 2 and r["nao_resolvidos"] == ["CidadeInexistenteZZZ"]
    assert "não resolv" in matriz_od.render_texto(r).lower()
    assert "não resolv" in matriz_od.construir_html(r).lower()
    # tudo resolvido → sem aviso de não resolvido.
    r2 = matriz_od.calcular("São Paulo/SP;Rio de Janeiro/RJ")
    assert not r2.get("nao_resolvidos") and "não resolv" not in matriz_od.render_texto(r2).lower()


# ============================================================================
#  CENTRO ÓTIMO — localização de facilidade (p-mediana / 1-centro).
# ============================================================================
def test_centro_otimo_estimado():
    from geo import centro_otimo as co
    # [PARIDADE-WEB] Sem OSRM, o padrão ESTIMA a distância rodoviária (sinuosidade do web). A ESCOLHA do
    # melhor candidato é a mesma da reta pura (Goiânia, central), mas os km ficam realistas.
    r = co.calcular("Goiânia/GO;São Paulo/SP;Recife/PE",
                    destinos="Belém/PA;Porto Alegre/RS;Rio de Janeiro/RJ;Manaus/AM")
    assert r["ok"] and r["modo"] == "estimado"
    assert len(r["candidatos"]) == 3 and len(r["destinos"]) == 4
    assert r["melhor_soma"]["nome"] == "Goiânia" and r["melhor_max"]["nome"] == "Goiânia"
    # soma/máx coerentes: o melhor por soma tem a menor soma do conjunto.
    menor_soma = min(c["soma_km"] for c in r["candidatos"])
    assert abs(r["melhor_soma"]["soma_km"] - menor_soma) < 0.5
    assert "CENTRO ÓTIMO" in co.render_texto(r)
    # paridade de TEMPO com o web: o pior caso ganha um tempo estimado de viagem (modelo 45/65 km/h).
    assert r["melhor_max"].get("max_min") is not None and r["melhor_max"]["max_min"] > 0
    # estimar=False → linha reta pura (modo 'geodesico'), MESMO vencedor, km menores (reta < viária).
    rg = co.calcular("Goiânia/GO;São Paulo/SP;Recife/PE",
                     destinos="Belém/PA;Porto Alegre/RS;Rio de Janeiro/RJ;Manaus/AM", estimar=False)
    assert rg["modo"] == "geodesico" and rg["melhor_soma"]["nome"] == "Goiânia"
    assert r["melhor_soma"]["soma_km"] > rg["melhor_soma"]["soma_km"]
    assert rg["melhor_max"].get("max_min") is None   # sem tempo a partir da reta pura (honesto)

def test_centro_otimo_self_dest():
    # Sem destinos → candidatos são os destinos (o mais central do conjunto).
    from geo import centro_otimo as co
    r = co.calcular("São Paulo/SP;Rio de Janeiro/RJ;Brasília/DF;Fortaleza/CE")
    assert r["ok"] and r.get("self_dest") and len(r["destinos"]) == 4
    # distância do candidato a si mesmo é 0 → entra na soma sem quebrar.
    assert all(c["soma_km"] >= 0 for c in r["candidatos"])

def test_centro_otimo_osrm_cai_em_estimado_sem_rede(monkeypatch):
    monkeypatch.setenv("OPENROTAS_NO_NET", "1")
    from geo import centro_otimo as co
    r = co.calcular("Goiânia/GO;São Paulo/SP", destinos="Belém/PA;Porto Alegre/RS", usar_osrm=True)
    assert r["ok"] and r["modo"] == "estimado"        # OSRM pulado sem rede → estimativa rodoviária

def test_centro_otimo_avisa_ambiguo_e_nao_resolvido():
    from geo import centro_otimo as co
    r = co.calcular("Bom Jesus;CidadeInexistenteZZZ;São Paulo/SP", destinos="Rio de Janeiro/RJ")
    assert r["ok"] and r["nao_resolvidos_candidatos"] == ["CidadeInexistenteZZZ"]
    t = co.render_texto(r); h = co.construir_html(r)
    assert "ambígua" in t and "não resolvido" in t.lower()
    assert "ambígua" in h and "não resolvido" in h.lower()

def test_centro_otimo_html_e_excel(tmp_path):
    import openpyxl
    from geo import centro_otimo as co
    r = co.calcular("Goiânia/GO;São Paulo/SP;Brasília/DF", destinos="Belém/PA;Porto Alegre/RS;Recife/PE")
    assert r["ok"]
    h = co.construir_html(r)
    assert h.startswith("<!doctype html>") and "Centro Ótimo" in h and "ork-theme-toggle" in h
    out = tmp_path / "c.html"
    assert co.gerar(out, r) == str(out) and out.exists()
    xo = tmp_path / "c.xlsx"
    assert co.exportar_excel(r, xo) == str(xo) and xo.exists()
    nomes = openpyxl.load_workbook(xo).sheetnames
    assert "Ranking" in nomes and "Distancias_km" in nomes

def test_centro_otimo_menos_de_1_destino():
    from geo import centro_otimo as co
    assert co.calcular("ZzzNadaQqq")["ok"] is False

def test_provisionamento_gera_centro_otimo(tmp_path, monkeypatch):
    import sys as _sys
    _sys.path.insert(0, "resources")
    import provisionamento as prov
    h = prov.gerar_centro_otimo("Goiânia/GO;São Paulo/SP", destinos="Belém/PA;Porto Alegre/RS")
    assert h and os.path.exists(h)
    assert prov.gerar_centro_otimo("ZzzCidadeInexistenteQqq") is None

def test_centro_k_hubs_exato():
    from geo import centro_otimo as co
    cands = "Goiânia/GO;São Paulo/SP;Recife/PE;Manaus/AM;Porto Alegre/RS"
    dests = "Belém/PA;Fortaleza/CE;Rio de Janeiro/RJ;Curitiba/PR;Salvador/BA;Brasília/DF;Manaus/AM"
    r = co.calcular_k(cands, destinos=dests, k=2)
    assert r["ok"] and r["k"] == 2 and r["metodo"] == "exato"
    assert len(r["solucao_soma"]["hubs"]) == 2 and len(r["solucao_max"]["hubs"]) == 2
    # cada destino atribuído exatamente uma vez
    assert len(r["solucao_soma"]["atribuicao"]) == 7
    # mais hubs não pioram a soma (k=2 <= k=1)
    r1 = co.calcular_k(cands, destinos=dests, k=1)
    assert r["solucao_soma"]["total_km"] <= r1["solucao_soma"]["total_km"] + 1e-6
    # p-centro otimiza o pior caso: seu máx <= máx da solução p-mediana
    assert r["solucao_max"]["max_km"] <= r["solucao_soma"]["max_km"] + 1e-6
    assert "HUBS" in co.render_texto(r) and "p-mediana" in co.render_texto(r)

def test_centro_k_cap_e_k_maior_que_candidatos():
    # k não pode exceder o nº de candidatos — é limitado a C (sem quebrar).
    from geo import centro_otimo as co
    r = co.calcular_k("São Paulo/SP;Rio de Janeiro/RJ", destinos="Curitiba/PR;Belém/PA", k=5)
    assert r["ok"] and r["k"] == 2 and len(r["solucao_soma"]["hubs"]) == 2

def test_centro_k_guloso_para_muitos_candidatos():
    # C > 22 → cai no método guloso (aproximação declarada), sem enumerar combinações demais.
    from geo import centro_otimo as co
    caps = ["São Paulo/SP", "Rio de Janeiro/RJ", "Belo Horizonte/MG", "Curitiba/PR", "Porto Alegre/RS",
            "Salvador/BA", "Recife/PE", "Fortaleza/CE", "Belém/PA", "Manaus/AM", "Goiânia/GO",
            "Brasília/DF", "Cuiabá/MT", "Campo Grande/MS", "Vitória/ES", "Natal/RN", "João Pessoa/PB",
            "Maceió/AL", "Aracaju/SE", "Teresina/PI", "São Luís/MA", "Palmas/TO", "Porto Velho/RO",
            "Rio Branco/AC", "Boa Vista/RR"]
    r = co.calcular_k(";".join(caps), destinos="Belém/PA;Rio de Janeiro/RJ;Fortaleza/CE;Manaus/AM", k=3)
    assert r["ok"] and r["metodo"] == "guloso" and len(r["solucao_soma"]["hubs"]) == 3

def test_centro_k_html_e_excel(tmp_path):
    import openpyxl
    from geo import centro_otimo as co
    r = co.calcular_k("Goiânia/GO;São Paulo/SP;Manaus/AM;Recife/PE",
                      destinos="Belém/PA;Rio de Janeiro/RJ;Fortaleza/CE;Curitiba/PR", k=2)
    assert r["ok"]
    h = co.construir_html(r)
    assert h.startswith("<!doctype html>") and "hubs" in h.lower() and "ork-theme-toggle" in h
    out = tmp_path / "k.html"
    assert co.gerar(out, r) == str(out) and out.exists()
    xo = tmp_path / "k.xlsx"
    assert co.exportar_excel(r, xo) == str(xo) and xo.exists()
    nomes = openpyxl.load_workbook(xo).sheetnames
    assert "Hubs" in nomes and "Atribuicao_pmediana" in nomes

def test_centro_k_via_wrapper(tmp_path):
    import sys as _sys
    _sys.path.insert(0, "resources")
    import provisionamento as prov
    h = prov.gerar_centro_otimo("Goiânia/GO;São Paulo/SP;Manaus/AM",
                                destinos="Belém/PA;Rio de Janeiro/RJ;Fortaleza/CE", k=2)
    assert h and os.path.exists(h)


# ============================================================================
#  ESTIMATIVA DE DISTÂNCIA RODOVIÁRIA (geo.distancia) — PARIDADE COM O WEB.
#  O desktop reportava a LINHA RETA crua no modo offline; o web nunca o faz —
#  ele estima a viária pela reta × fator de sinuosidade (obter_fator_desvio_
#  rodoviario). Este módulo traz o MESMO fator ao desktop. Testes PUROS.
# ============================================================================
def test_distancia_fator_paridade_com_web():
    # Fatores IDÊNTICOS ao web: 1,45 (<5), 1,35 (<20), 1,25 (<100), 1,18 (>=100).
    from geo import distancia as d
    assert d.fator_desvio(4.9) == 1.45 and d.fator_desvio(5.0) == 1.35
    assert d.fator_desvio(19.9) == 1.35 and d.fator_desvio(20.0) == 1.25
    assert d.fator_desvio(99.9) == 1.25 and d.fator_desvio(100.0) == 1.18
    assert d.fator_desvio(5000.0) == 1.18
    assert d.fator_desvio("lixo") == 1.18            # defensivo: cai no fator conservador de longa

def test_distancia_estimativa_escalar_e_piso():
    from geo import distancia as d
    # reta × fator, com a reta como PISO (fator >= 1 → estimativa nunca menor que a reta).
    assert round(d.estimar_rodoviaria_km(100.0), 2) == 118.0
    assert round(d.estimar_rodoviaria_km(10.0), 2) == 13.5
    for r in (1.0, 7.0, 30.0, 120.0, 800.0):
        assert d.estimar_rodoviaria_km(r) >= r        # nenhuma estrada é mais curta que a reta
    assert d.estimar_rodoviaria_km(0) == 0.0 and d.estimar_rodoviaria_km(-3) == 0.0
    assert d.estimar_rodoviaria_km(None) == 0.0       # defensivo

def test_distancia_vetor_bate_escalar():
    import numpy as np
    from geo import distancia as d
    retas = [0.0, 3.0, 10.0, 19.9, 20.0, 50.0, 99.9, 100.0, 400.0, -1.0]
    v = d.estimar_rodoviaria_vetor(retas)
    esc = [d.estimar_rodoviaria_km(x) for x in retas]
    assert np.allclose(np.asarray(v, dtype="float64"), esc)   # vetorizado == escalar, faixa a faixa

def test_distancia_tempo_modelo_web():
    from geo import distancia as d
    # 45 km/h abaixo de 50 km, 65 km/h acima (modelo do web).
    assert d.estimar_tempo_min(45.0) == 60.0          # 45 km a 45 km/h = 60 min
    assert d.estimar_tempo_min(65.0) == 60.0          # 65 km a 65 km/h = 60 min
    assert d.estimar_tempo_min(0) == 0.0 and d.estimar_tempo_min(None) == 0.0

def test_distancia_rotulos_de_modo():
    from geo import distancia as d
    assert d.rotulo_modo("estimado") == "estimado"
    assert d.rotulo_modo("rodoviario") == "rodoviário"
    assert d.rotulo_modo("geodesico") == "geodésico"
    assert "sinuosidade" in d.descricao_modo("estimado")
    assert "OSRM" in d.descricao_modo("rodoviario")

def test_distancia_haversine_vs_planar():
    # Haversine honesta: SP(-46.63,-23.55)→RJ(-43.17,-22.91) ~ 360 km em linha reta (não a planar inflada).
    import numpy as np
    from geo import distancia as d
    km = float(d.haversine_km(-46.63, -23.55, -43.17, -22.91))
    assert 350 <= km <= 375, km
    # vetor: 1º elemento == escalar; o mesmo ponto dá 0.
    v = d.haversine_km(-46.63, -23.55, np.array([-43.17, -46.63]), np.array([-22.91, -23.55]))
    assert abs(float(v[0]) - km) < 0.5 and float(v[1]) < 0.1
    # a planar (dlon²+dlat²)×111 INFLA a leste-oeste: a haversine é menor para um par E-W puro no Sul.
    import math
    lo0, la0, lo1, la1 = -51.0, -30.0, -48.0, -30.0   # 3° de longitude puros a 30°S
    planar = math.hypot(lo1 - lo0, la1 - la0) * 111.0
    hv = float(d.haversine_km(lo0, la0, lo1, la1))
    assert hv < planar and hv == __import__("pytest").approx(3 * 111.0 * math.cos(math.radians(30)), rel=0.03)

def test_distancia_valida_coordenada_br():
    from geo import distancia as d
    assert d.validar_coordenada_br(-46.63, -23.55) == []          # SP, plausível
    assert d.validar_coordenada_br(-47.88, -15.79) == []          # Brasília, plausível
    # troca lat/lon (usuário digitou lat,lon): lon=-23.55 (|<34|), lat=-46.63 (|>=34|) → sinalizado
    assert d.validar_coordenada_br(-23.55, -46.63)
    # fora do Brasil (Lisboa ~ -9.14, 38.7): |lat|>=34 e fora dos limites
    assert d.validar_coordenada_br(-9.14, 38.7)
    assert d.validar_coordenada_br(0.0, 0.0)                      # (0,0) ausente
    assert d.validar_coordenada_br("x", "y")                      # não-numérico

def test_matriz_sinaliza_coordenada_suspeita():
    # Uma coordenada CRUA trocada (lat,lon) entra na matriz mas é SINALIZADA (não silenciada).
    from geo import matriz_od
    r = matriz_od.calcular([(-46.63, -23.55), (-23.55, -46.63)])   # 2º ponto está com lon,lat trocados
    assert r["ok"] and r["resumo"]["coord_suspeitas"] == 1
    assert any(p.get("coord_suspeita") for p in r["pontos"])
    assert "suspeita" in matriz_od.render_texto(r).lower()
    assert "suspeita" in matriz_od.construir_html(r).lower()

def test_isocronas_sinaliza_origem_suspeita():
    from geo import isocronas
    r = isocronas.calcular("-23.55,-46.63", faixas_km=(100, 300))   # lon,lat trocados na origem
    assert r["ok"] and r["origem"].get("coord_suspeita")
    assert "suspeita" in isocronas.render_texto(r).lower()
    assert "suspeita" in isocronas.construir_html(r).lower()

def test_distancia_formatar_duracao():
    from geo import distancia as d
    assert d.formatar_duracao(0) == "" and d.formatar_duracao(None) == ""
    assert d.formatar_duracao(45) == "45 min"
    assert d.formatar_duracao(60) == "1h00min"
    assert d.formatar_duracao(125) == "2h05min"

def test_matriz_tempo_no_par_distante():
    # [PARIDADE-WEB · TEMPO] No modo estimado/rodoviário o par mais distante ganha tempo estimado; na
    # linha reta pura (geodesico) não há tempo (não se estima viagem a partir da reta).
    from geo import matriz_od
    r = matriz_od.calcular("São Paulo/SP;Rio de Janeiro/RJ;Brasília/DF")       # default: estimado
    assert r["modo"] == "estimado" and r["resumo"]["par_distante"].get("min") is not None
    assert "≈" in matriz_od.render_texto(r) and "≈" in matriz_od.construir_html(r)
    rg = matriz_od.calcular("São Paulo/SP;Rio de Janeiro/RJ;Brasília/DF", estimar=False)
    assert rg["modo"] == "geodesico" and rg["resumo"]["par_distante"].get("min") is None

def test_isocronas_tempo_por_faixa():
    from geo import isocronas
    r = isocronas.calcular("Brasília/DF", faixas_km=(100, 300, 600))           # default: estimado
    assert r["modo"] == "estimado"
    assert all(b.get("tempo_min") and b["tempo_min"] > 0 for b in r["bandas"])
    assert "≈" in isocronas.render_texto(r)
    rg = isocronas.calcular("Brasília/DF", faixas_km=(100, 300, 600), estimar=False)
    assert all(b.get("tempo_min") is None for b in rg["bandas"])               # reta pura → sem tempo

def test_repo_proximidade_corrige_longitude(repo):
    # A correção cos(lat) AMPLIA a caixa a leste-oeste: num ponto do SUL, proximidade deve cobrir
    # >= o que uma caixa de grau UNIFORME (sem a correção) cobriria. Prova a melhoria sem contagem fixa.
    import math
    from geo import repositorio
    lon, lat, raio = -51.23, -30.03, 40.0        # Porto Alegre/RS (cos(30) ~ 0.87 → correção relevante)
    r = repo.proximidade("drenagem", lon, lat, raio_km=raio)
    rg = max(0.0, raio) / repositorio._GRAU_KM
    naive = repo.intersecta_bbox("drenagem", lon - rg, lat - rg, lon + rg, lat + rg)   # caixa uniforme
    assert r["total"] >= int(naive.size)         # a corrigida nunca cobre menos que a uniforme


# ============================================================================
#  CONSCIÊNCIA FLUVIAL — municípios de acesso só fluvial (rio) recebem o fator de
#  sinuosidade FLUVIAL (~3×), não o rodoviário. Lista oficial IBGE REGIC (17 rios
#  amazônicos; Fernando de Noronha, ilha oceânica, é EXCLUÍDO do ~3×). Paridade web.
# ============================================================================
def test_distancia_lista_fluvial_oficial():
    from geo import distancia as d
    assert d.municipio_acesso_fluvial("1300201") is True       # Atalaia do Norte/AM (rio)
    assert d.municipio_acesso_fluvial("1302108") is True       # Japurá/AM (rio)
    assert d.municipio_acesso_fluvial("2605459") is False      # Fernando de Noronha — ilha oceânica, NÃO rio
    assert d.municipio_acesso_fluvial("3550308") is False      # São Paulo — rodoviário
    assert d.municipio_acesso_fluvial("1.300.201") is True     # normaliza pontuação/zeros
    assert d.municipio_acesso_fluvial("") is False and d.municipio_acesso_fluvial(None) is False
    assert len(d._MUNICIPIOS_ACESSO_FLUVIAL) == 17

def test_distancia_fator_fluvial():
    from geo import distancia as d
    # fluvial = reta × ~3 (muito maior que o rodoviário ~1,2); piso continua sendo a reta.
    assert d.estimar_rodoviaria_km(100.0, fluvial=True) == 300.0
    assert d.estimar_rodoviaria_km(100.0, fluvial=True) > d.estimar_rodoviaria_km(100.0)
    assert d.estimar_rodoviaria_km(100.0, fluvial=True) >= 100.0
    # vetor com máscara por elemento: só o 2º é fluvial.
    import numpy as np
    v = d.estimar_rodoviaria_vetor([100.0, 100.0], fluvial=[False, True])
    assert abs(float(v[0]) - 118.0) < 0.1 and abs(float(v[1]) - 300.0) < 0.1
    # máscara a partir de Códigos IBGE.
    assert d.mascara_acesso_fluvial(["1300201", "3550308"]) == [True, False]
    # tempo fluvial é mais lento (~22 km/h) que o rodoviário.
    assert d.estimar_tempo_min(100.0, fluvial=True) > d.estimar_tempo_min(100.0)

def test_matriz_fluvial_aplica_fator_e_sinaliza():
    # Manaus ↔ Atalaia do Norte (acesso só por rio): a estimativa usa o fator FLUVIAL (~3× a reta) e o
    # relatório sinaliza. Fernando de Noronha (ilha oceânica) NÃO entra no ~3× nem é marcado fluvial.
    from geo import matriz_od
    m = matriz_od.calcular("Manaus/AM;Atalaia do Norte/AM")
    mg = matriz_od.calcular("Manaus/AM;Atalaia do Norte/AM", estimar=False)
    assert m["resumo"]["acesso_fluvial"] == 1
    assert abs(m["matriz"][0][1] / mg["matriz"][0][1] - 3.0) < 0.05      # ~3× a reta
    assert "🛶" in matriz_od.render_texto(m) and "🛶" in matriz_od.construir_html(m)
    f = matriz_od.calcular("Recife/PE;Fernando de Noronha/PE")
    assert f["resumo"]["acesso_fluvial"] == 0                            # ilha oceânica não é "fluvial"
    assert f["matriz"][0][1] / matriz_od.calcular("Recife/PE;Fernando de Noronha/PE",
                                                  estimar=False)["matriz"][0][1] < 1.5   # fator rodoviário

def test_centro_osrm_celula_sem_rota_cai_na_estimativa(monkeypatch):
    # [OSRM-PARCIAL] No centro ótimo, se o /table não roteia UM par (ilha/fluvial), antes a matriz INTEIRA
    # era descartada (caía tudo na estimativa). Agora usa o OSRM real onde há e estima só a célula sem rota.
    from geo import dossie_rota as dr, centro_otimo as co
    monkeypatch.setattr(dr, "osrm_ativo", lambda *a, **k: "http://localhost:5000")
    # _matriz_osrm_cd dublado: 2 candidatos × 2 destinos; um par (cand0→dest1) sem rota (None).
    monkeypatch.setattr(co, "_matriz_osrm_cd", lambda cand, dest, url: [[100.0, None], [120.0, 90.0]])
    r = co.calcular("Manaus/AM;Belém/PA", destinos="Santarém/PA;Macapá/AP")
    assert r["ok"] and r["modo"] == "rodoviario-parcial"       # real onde há + estimativa na célula sem rota
    c0 = next(c for c in r["candidatos"] if c["distancias"][0] == 100.0)  # usou o OSRM real (100)
    assert c0["distancias"][1] > 0                             # a célula None virou estimativa (não 0/reta crua)

def test_centro_fluvial_sinaliza():
    from geo import centro_otimo as co
    r = co.calcular("Manaus/AM;Tefé/AM", destinos="Atalaia do Norte/AM;Tabatinga/AM")
    assert any(d.get("acesso_fluvial") for d in r["destinos"]) or any(c.get("acesso_fluvial") for c in r["candidatos"])
    assert "🛶" in co.render_texto(r) and "🛶" in co.construir_html(r)

def test_isocronas_fluvial_recebe_fator():
    # De Manaus, as faixas alcançam municípios fluviais; eles recebem o ~3× e são contados/sinalizados.
    from geo import isocronas
    r = isocronas.calcular("Manaus/AM", faixas_km=(300, 700, 1500, 3000))
    assert "acesso_fluvial_alcancados" in r
    if r["acesso_fluvial_alcancados"] > 0:                               # dependente da base, mas tipicamente >0
        assert "🛶" in isocronas.render_texto(r)


# ============================================================================
#  SENSOR DE CIRCUIDADE (viária ÷ reta) no dossiê — só SINALIZA barreira/implausível. Paridade web.
# ============================================================================
def test_circuidade_sensor_bandas():
    from geo import distancia as d
    assert d.circuidade_suspeita(2.5, 50) == (True, "PROVAVEL_BARREIRA")   # >= 2,0 em qualquer distância
    assert d.circuidade_suspeita(1.9, 10) == (False, "")                   # curto: gatilho 2,2 → não
    assert d.circuidade_suspeita(1.9, 30) == (True, "INVESTIGAR")          # médio: gatilho 1,8 → investigar
    assert d.circuidade_suspeita(1.2, 200) == (False, "")                  # longo: gatilho 1,6 → não
    assert d.circuidade_suspeita("x", 10) == (False, "")                   # defensivo
    assert d.circuidade_suspeita(1.0, 0) == (False, "")                    # reta nula → nada

def test_indice_dor_logistica_paridade_web():
    from geo import distancia as d
    # Mesmos pesos/gatilhos do web: distância 30% (ref 1500), sinuosidade 25% (1,4→3,0×), balsa 25%,
    # isolamento 20%. Caso saturado: dv=1000/lr=300 (sin satura 100), com balsa e fluvial.
    assert d.indice_dor_logistica(1000, 300, tem_balsa=True, fluvial=True) == 90.0
    # Rodoviário curto sem complicação → dor baixa (só distância e sinuosidade leve).
    assert d.indice_dor_logistica(100, 70) == 2.4
    # Balsa sozinha pesa 25 pontos; isolamento sozinho pesa 20.
    assert d.indice_dor_logistica(0, 0, tem_balsa=True) == 25.0
    assert d.indice_dor_logistica(0, 0, fluvial=True) == 20.0
    # Defensivo: entradas inválidas não levantam e contribuem 0 no seu eixo.
    assert d.indice_dor_logistica(None, None) == 0.0
    assert d.indice_dor_logistica("x", 70) == 0.0
    # Clampeado em [0, 100].
    assert 0.0 <= d.indice_dor_logistica(99999, 1, tem_balsa=True, fluvial=True) <= 100.0

def test_dor_logistica_nivel():
    from geo import distancia as d
    assert d.dor_logistica_nivel(10) == "BAIXA"
    assert d.dor_logistica_nivel(50) == "MÉDIA"
    assert d.dor_logistica_nivel(90) == "ALTA"
    assert d.dor_logistica_nivel(67) == "ALTA" and d.dor_logistica_nivel(34) == "MÉDIA"
    assert d.dor_logistica_nivel("x") == ""

def test_dossie_circuidade_rodoviaria_normal():
    # Par rodoviário normal (SP→RJ): circuidade ~1,2× (fator de estimativa), SEM suspeita de barreira.
    from geo import dossie_rota
    dd, _ = dossie_rota.dossie_por_nomes("São Paulo/SP;Rio de Janeiro/RJ")
    c = dd.get("circuidade")
    assert c and c["fator"] and 1.1 <= c["fator"] <= 1.5
    assert c["suspeita"] is False and c["nivel"] == ""
    assert "Circuidade" in dossie_rota.render_texto(dd)

def test_dossie_circuidade_fluvial_sinaliza_barreira():
    # Par de acesso só fluvial (Manaus→Atalaia do Norte): a estimativa fluvial (~3×) dispara PROVÁVEL
    # BARREIRA — o sensor corrobora a dependência de rio, mesmo offline (sem OSRM).
    from geo import dossie_rota
    dd, _ = dossie_rota.dossie_por_nomes("Manaus/AM;Atalaia do Norte/AM")
    c = dd.get("circuidade")
    assert c and c["fator"] >= 2.0 and c["nivel"] == "PROVAVEL_BARREIRA" and c["suspeita"] is True
    assert c["fluvial"] is True
    assert any("BARREIRA" in a for a in dd["alertas"])
    assert "provável barreira" in dossie_rota.render_texto(dd)

def test_dossie_dor_logistica_presente_e_coerente():
    # O dossiê expõe a 'dor logística' (0-100) a partir dos sinais já medidos. Par rodoviário curto e
    # limpo (Campinas→Ribeirão Preto, sem travessia no corredor nem isolamento) → dor BAIXA, sem balsa
    # nem fluvial; o render de texto mostra a linha. (balsa reflete o MESMO sinal de travessia no
    # corredor que o dossiê já usa no alerta 'possível dependência aquática'.)
    from geo import dossie_rota
    dd, _ = dossie_rota.dossie_por_nomes("Campinas/SP;Ribeirão Preto/SP")
    dl = dd.get("dor_logistica")
    assert dl and 0.0 <= dl["indice"] <= 100.0
    assert dl["nivel"] in ("BAIXA", "MÉDIA", "ALTA")
    assert dl["balsa"] is False and dl["fluvial"] is False
    assert dl["nivel"] == "BAIXA"                       # par tronco rodoviário não é destino "difícil"
    assert "Dor logística" in dossie_rota.render_texto(dd)

def test_dossie_dor_logistica_fluvial_alta():
    # Par de acesso só fluvial (Manaus→Atalaia do Norte): distância longa + sinuosidade ~3× + isolamento
    # REGIC → dor ALTA, com alerta correspondente. Mede dificuldade (não decide rota), mesmo offline.
    from geo import dossie_rota
    dd, _ = dossie_rota.dossie_por_nomes("Manaus/AM;Atalaia do Norte/AM")
    dl = dd.get("dor_logistica")
    assert dl and dl["fluvial"] is True and dl["nivel"] == "ALTA" and dl["indice"] >= 67.0
    assert any("Dor logística ALTA" in a for a in dd["alertas"])


# ============================================================================
#  NAVEGABILIDADE POR NOME — supre a lacuna 'desconhecido' (~86%) da base de drenagem
#  reconhecendo os grandes rios navegáveis pelo nome. Precisão (sem falso-positivo).
# ============================================================================
def test_rio_navegavel_por_nome_precisao():
    from geo import distancia as d
    for nm in ("Rio Amazonas", "Solimões", "Rio Negro", "Rio São Francisco", "Rio Paraná", "Rio Madeira"):
        assert d.rio_navegavel_por_nome(nm) is True, nm
    for nm in ("Laguna dos Patos", "Córrego do Meio", "Rio Grande", "Igarapé Una", "Riacho Seco", "", None,
               "Rio São Francisco Falso", "Rio São Francisco Verdadeiro"):   # qualificador no fim = OUTRO rio
        assert d.rio_navegavel_por_nome(nm) is False, nm   # fronteira de palavra: 'una' não casa 'laguna'

def test_barragens_do_rio_ressalva():
    # Rios navegáveis com barragem MAIOR que interrompe a navegação de passagem (ressalva honesta).
    from geo import distancia as d
    assert d.barragens_do_rio("Rio Paraná") == ["Itaipu"]
    assert "Belo Monte" in d.barragens_do_rio("Rio Xingu")
    assert d.barragens_do_rio("Rio São Francisco")                 # Sobradinho/Três Marias/Xingó
    assert d.barragens_do_rio("Rio Amazonas") == []                # navegável e sem barragem
    assert d.barragens_do_rio("Laguna") == [] and d.barragens_do_rio("") == []
    assert d.barragens_do_rio("Rio São Francisco Falso") == []     # outro rio — não herda a ressalva

def test_dossie_ressalva_barragem_em_rio_navegavel():
    # Dossiê que cruza o Paraná: reconhece navegável E a ressalva de Itaipu (barragem que bloqueia).
    from geo import dossie_rota
    dd, _ = dossie_rota.dossie_por_nomes("Foz do Iguaçu/PR;Guaíra/PR")   # corredor sobre o Rio Paraná
    h = dd.get("hidrografia", {})
    rb = h.get("rios_navegaveis_com_barragem", [])
    if rb:                                                          # dependente de o corredor pegar o Paraná
        assert any("Itaipu" in x["barragens"] for x in rb)
        assert "barragem" in dossie_rota.render_texto(dd).lower()

def test_dossie_navegabilidade_por_nome_supre_base():
    # Manaus→Parintins cruza o Amazonas; a base marca 'desconhecido' (navegaveis=0), mas o
    # reconhecimento por nome recupera o Rio Amazonas como navegável.
    from geo import dossie_rota
    dd, _ = dossie_rota.dossie_por_nomes("Manaus/AM;Parintins/AM")
    h = dd.get("hidrografia", {})
    assert h.get("navegaveis_nomeados", 0) >= 1
    assert any("Amazonas" in n for n in h.get("rios_navegaveis_nome", []))
    # cada rio nomeado carrega a tag navegavel; e o texto expõe a lista por nome.
    assert all("navegavel" in r for r in h.get("rios_nomeados", []))
    assert "navegáveis reconhecidos por nome" in dossie_rota.render_texto(dd)


# ============================================================================
#  QUALIDADE DE COORDENADAS — guarda anti-regressão do defeito Brasília/Vitória.
# ============================================================================
def test_coverage_ponto_no_corpo_helper():
    import struct
    from audit import coverage_auditor as ca
    ring = [(-47, -16), (-47, -15), (-46, -15), (-46, -16), (-47, -16)]
    wkb = struct.pack(">BII", 0, 3, 1) + struct.pack(">I", len(ring)) + \
        b"".join(struct.pack(">dd", x, y) for x, y in ring)
    assert ca.ponto_no_corpo(-46.5, -15.5, wkb) is True       # dentro do corpo
    assert ca.ponto_no_corpo(-46.05, -15.5, wkb) is True      # perto da borda
    assert ca.ponto_no_corpo(-33.8, -3.9, wkb) is False       # no oceano, longe → suspeito
    assert ca.ponto_no_corpo(-46.5, -15.5, b"lixo") is True   # WKB inválido → não julga

def test_coverage_auditar_coordenadas_limpo_pos_correcao():
    from audit import coverage_auditor as ca
    r = ca.auditar_coordenadas()
    # após a correção de Brasília/Vitória, nenhum município deve estar fora do corpo principal
    assert r["total"] > 5000 and r["status"] == "OK" and r["suspeitos"] == []

def test_integridade_reporta_coordenadas():
    from audit import integridade as ig
    aud = ig.auditar()
    assert "coordenadas" in aud and aud["resumo"].get("coordenadas_suspeitas") == 0
    assert aud["coordenadas"].get("checado") is True
    assert "Coordenadas municipais" in ig.render_texto(aud)

def test_integridade_coordenadas_opcional():
    # coordenadas=False pula a varredura pesada (usado pela Saúde Nacional) e NÃO alega ter checado
    from audit import integridade as ig
    r = ig.auditar(coordenadas=False)
    assert r["coordenadas"]["checado"] is False and r["resumo"]["coordenadas_suspeitas"] == 0
    assert "Coordenadas municipais" not in ig.render_texto(r)


# ============================================================================
#  BOTÕES DA CENTRAL DE DADOS — wrappers que geram cada análise avançada (HTML).
# ============================================================================
def test_prov_wrappers_analises_avancadas(monkeypatch, tmp_path):
    import sys as _sys, pathlib as _pl
    _sys.path.insert(0, str(_pl.Path(__file__).resolve().parents[1]))
    monkeypatch.setattr(prov, "pasta_por_tipo", lambda tipo: tmp_path / tipo)
    for fn in (prov.gerar_painel_executivo, prov.gerar_mapa_cobertura,
               prov.gerar_conectividade_multimodal):
        p = fn()
        assert p and _pl.Path(p).exists() and _pl.Path(p).read_text(encoding="utf-8").startswith("<!doctype html>")
    pi = prov.gerar_isocronas("São Paulo/SP")
    assert pi and _pl.Path(pi).exists()
    pm = prov.gerar_matriz_od("São Paulo/SP;Rio de Janeiro/RJ;Brasília/DF")
    assert pm and _pl.Path(pm).exists()
    assert prov.gerar_isocronas("ZzzCidadeInexistenteQqq") is None      # origem não resolve → None


# ============================================================================
#  DOSSIÊ POR NOME DE CIDADE (desktop) — resolução via índice de municípios IBGE.
# ============================================================================
def test_resolver_local_parse():
    from geo import dossie_rota as dr
    assert dr._parse_local("São Paulo/SP") == ("sao paulo", "SP")
    assert dr._parse_local("Rio de Janeiro, RJ") == ("rio de janeiro", "RJ")
    assert dr._parse_local("Belo Horizonte - MG") == ("belo horizonte", "MG")
    assert dr._parse_local("Curitiba PR") == ("curitiba", "PR")
    assert dr._parse_local("Manaus") == ("manaus", None)

def test_resolver_local_cidade_real():
    from geo import dossie_rota as dr
    r = dr.resolver_local("São Paulo/SP")
    assert r["ok"] and r["uf"] == "SP" and r["nome"].lower().startswith("são paulo")
    assert -47 < r["lon"] < -46 and -24 < r["lat"] < -23

def test_resolver_local_desambigua_por_uf():
    from geo import dossie_rota as dr
    # "Bom Jesus" existe em várias UFs; com /PI deve resolver no Piauí
    r = dr.resolver_local("Bom Jesus/PI")
    assert r["ok"] and r["uf"] == "PI"

def test_resolver_local_inexistente():
    from geo import dossie_rota as dr
    r = dr.resolver_local("Cidade Inexistente XYZ")
    assert r["ok"] is False

def test_dossie_por_nomes_sp_rio():
    from geo import dossie_rota as dr
    dd, resol = dr.dossie_por_nomes("São Paulo/SP;Rio de Janeiro/RJ")
    assert not dd.get("erro")
    assert dd["origem"]["municipio"]["uf"] == "SP" and dd["destino"]["municipio"]["uf"] == "RJ"
    assert dd["hidrografia"]["trechos"] > 100
    assert len(dd["locais"]) == 2 and all(l["ok"] for l in dd["locais"])

def test_dossie_por_nomes_nenhuma_resolvida():
    from geo import dossie_rota as dr
    dd, resol = dr.dossie_por_nomes("Xyzabc123;Qwerty999")
    assert dd.get("erro") and all(not r["ok"] for r in resol)

def test_dossie_por_nomes_cidade_ambigua_avisa(tmp_path):
    # Mesma honestidade da matriz/isócronas: 'Bom Jesus' (homônima em várias UFs) é resolvida pela
    # 1ª opção, mas o RELATÓRIO (texto e HTML, usados pela GUI e pelos exports) DEVE avisar — não só
    # a CLI. Dossiê por coordenadas e nomes com /UF não geram aviso.
    from geo import dossie_rota as dr
    dd, _ = dr.dossie_por_nomes("Bom Jesus;São Paulo/SP")
    assert not dd.get("erro")
    amb = dr._locais_ambiguos(dd)
    assert len(amb) == 1 and "Bom Jesus" in amb[0]
    assert "⚠" in dr.render_texto(dd)
    h = dr.gerar_html(dd, tmp_path / "d.html")
    assert h and "ambígua" in open(h, encoding="utf-8").read()
    # com /UF explícita → sem ambiguidade nem aviso
    dd2, _ = dr.dossie_por_nomes("Bom Jesus/PI;São Paulo/SP")
    assert dr._locais_ambiguos(dd2) == [] and "⚠ Cidade" not in dr.render_texto(dd2)

def test_dossie_por_nomes_local_nao_resolvido_avisa(tmp_path):
    # Um nome não encontrado fica de fora do dossiê — relatório (texto e HTML) deve avisar, não só a CLI.
    from geo import dossie_rota as dr
    dd, _ = dr.dossie_por_nomes("São Paulo/SP;CidadeInexistenteZZZ;Rio de Janeiro/RJ")
    assert not dd.get("erro")
    assert dr._locais_nao_resolvidos(dd) == ["CidadeInexistenteZZZ"]
    assert "não resolv" in dr.render_texto(dd).lower()
    h = dr.gerar_html(dd, tmp_path / "d.html")
    assert h and "não resolv" in open(h, encoding="utf-8").read().lower()
    # todas resolvidas → sem aviso
    dd2, _ = dr.dossie_por_nomes("São Paulo/SP;Rio de Janeiro/RJ")
    assert dr._locais_nao_resolvidos(dd2) == [] and "não resolv" not in dr.render_texto(dd2).lower()


# ============================================================================
#  ROTA REAL OSRM + LOTE + SOBREPOSIÇÃO NO MAPA (desktop).
# ============================================================================
def test_osrm_parse_geojson():
    from geo import dossie_rota as dr
    j = '{"code":"Ok","routes":[{"distance":1000,"duration":60,"geometry":{"type":"LineString",'\
        '"coordinates":[[-46.6,-23.5],[-46.0,-23.2],[-43.2,-22.9]]}}]}'
    coords, dist, dur = dr._parse_osrm_geojson(j)
    assert coords == [(-46.6, -23.5), (-46.0, -23.2), (-43.2, -22.9)] and dist == 1000
    assert dr._parse_osrm_geojson('{"code":"NoRoute"}') == (None, None, None)

def test_osrm_no_net_retorna_none(monkeypatch):
    from geo import dossie_rota as dr
    monkeypatch.setenv("OPENROTAS_NO_NET", "1")
    assert dr.geometria_osrm([(-46.6, -23.5), (-43.2, -22.9)], url="http://x") is None

def test_dossie_usa_rota_real_quando_osrm_responde(monkeypatch):
    from geo import dossie_rota as dr
    # polilinha "real" densa (simula OSRM), hugging um caminho diferente da reta
    fake = [(-46.63, -23.55), (-46.0, -23.2), (-44.5, -23.0), (-43.20, -22.90)]
    # a assinatura ganhou retornar_distancia → o dublê devolve (linha, dist_m, dur_s) quando pedido.
    monkeypatch.setattr(dr, "geometria_osrm",
                        lambda coords, url=None, retornar_distancia=False:
                        (fake, 430000.0, 25200.0) if retornar_distancia else fake)
    dd = dr.dossie([(-46.63, -23.55), (-43.20, -22.90)], usar_osrm=True)
    assert dd["rota_real"] is True and dd["rota_pontos_osrm"] == 4
    # a circuidade agora usa a viária REAL do OSRM (430 km) vs a reta (~360 km) → ~1,2×, sem barreira.
    c = dd.get("circuidade")
    assert c and c["fonte"].startswith("OSRM") and 1.1 <= c["fator"] <= 1.4 and c["suspeita"] is False

def test_dossie_sem_osrm_degrada(monkeypatch):
    from geo import dossie_rota as dr
    monkeypatch.setattr(dr, "geometria_osrm",
                        lambda coords, url=None, retornar_distancia=False:
                        (None, None, None) if retornar_distancia else None)
    dd = dr.dossie([(-46.63, -23.55), (-43.20, -22.90)], usar_osrm=True)
    assert dd["rota_real"] is False     # sem OSRM → corredor reta, honesto
    # mesmo sem OSRM, a circuidade existe via ESTIMATIVA (fonte declarada), sem suspeita p/ par normal.
    c = dd.get("circuidade")
    assert c and "estimativa" in c["fonte"] and c["suspeita"] is False

def test_lote_ler_pares_e_processar(tmp_path):
    import pandas as pd
    from geo import dossie_lote as dl
    ent = tmp_path / "pares.xlsx"
    pd.DataFrame({"origem": ["São Paulo/SP", "Curitiba/PR"],
                  "destino": ["Rio de Janeiro/RJ", "Florianópolis/SC"]}).to_excel(ent, index=False)
    assert dl.ler_pares(ent) == [("São Paulo/SP", "Rio de Janeiro/RJ"), ("Curitiba/PR", "Florianópolis/SC")]
    out = tmp_path / "consol.xlsx"
    r = dl.processar_lote(ent, out)
    assert r["ok"] and r["linhas"] == 2 and os.path.exists(r["saida"])
    import openpyxl
    wb = openpyxl.load_workbook(out)
    assert "Rotas" in wb.sheetnames
    ws = wb["Rotas"]
    assert ws.max_row == 3        # cabeçalho + 2 rotas

def test_lote_coords_e_sem_pares(tmp_path):
    from geo import dossie_lote as dl
    assert dl._coords_de("-46.6,-23.5") == (-46.6, -23.5)
    assert dl._coords_de("São Paulo") is None
    r = dl.processar_lote(tmp_path / "naoexiste.xlsx")
    assert r["ok"] is False

def test_mapa_cobertura_sobrepoe_rota():
    from audit import coverage_map as cm
    h = cm.construir_html(rapido=True, rota="São Paulo/SP;Rio de Janeiro/RJ")
    assert "Rota sobreposta" in h and "<th>Rota</th>" in h
    assert "●" in h             # UFs atravessadas marcadas

def test_mapa_cobertura_sem_rota_igual_antes():
    from audit import coverage_map as cm
    h = cm.construir_html(rapido=True)
    assert "Rota sobreposta" not in h and "<th>Rota</th>" not in h


# ============================================================================
#  SAÚDE DOS DADOS NACIONAIS (§39) — painel único que consolida os auditores.
# ============================================================================
@pytest.fixture(scope="module")
def saude():
    from audit import saude_nacional as sn
    return sn.diagnostico(rapido=True)

def test_saude_dimensoes_e_resumo(saude):
    nomes = {x["nome"] for x in saude["dimensoes"]}
    assert "Cobertura territorial" in nomes and "Integridade geométrica" in nomes
    assert "Rede multimodal" in nomes and "Prontidão offline" in nomes
    r = saude["resumo"]
    assert r["total"] == len(saude["dimensoes"]) == 5
    assert r["ok"] + r["parciais"] + r["ausentes"] == r["total"]
    assert r["status"] in ("OK", "PARCIAL", "AUSENTE")

def test_saude_cobertura_ok(saude):
    cob = next(x for x in saude["dimensoes"] if x["nome"] == "Cobertura territorial")
    assert cob["status"] == "OK" and "27/27" in cob["resumo"]
    integ = next(x for x in saude["dimensoes"] if x["nome"] == "Integridade geométrica")
    assert integ["status"] == "OK"

def test_saude_offline_honesta(saude):
    # Prontidão offline honesta em QUALQUER ambiente: OK quando tudo (grafo incluído) está pronto,
    # PARCIAL caso contrário — e, quando PARCIAL, o motivo é explicado (menção ao grafo). Testar a
    # invariante torna o teste portável (alguns ambientes já têm o grafo instalado) sem perder o que
    # importa: um estado PARCIAL nunca é silencioso.
    off = next(x for x in saude["dimensoes"] if x["nome"] == "Prontidão offline")
    assert off["status"] in ("OK", "PARCIAL")
    if off["status"] == "PARCIAL":
        assert "grafo" in (off["resumo"] + off["detalhe"]).lower()

def test_saude_render_e_html(saude):
    from audit import saude_nacional as sn
    txt = sn.render_texto(saude)
    assert "SAÚDE DOS DADOS NACIONAIS" in txt and "Resumo:" in txt
    h = sn.construir_html(rapido=True)
    assert "<!doctype html>" in h and "Saúde dos Dados Nacionais" in h
    assert "prefers-color-scheme" in h and "Ver detalhes" in h

def test_saude_gera_arquivo(tmp_path):
    from audit import saude_nacional as sn
    out = tmp_path / "saude.html"
    got = sn.gerar(out, rapido=True)
    assert got == str(out) and out.exists()
    assert out.read_text(encoding="utf-8").startswith("<!doctype html>")


# ============================================================================
#  DOWNLOADER ROBUSTO (§18/§20) — retomada, velocidade/ETA, retry, sha256.
#  Opener INJETÁVEL → testável sem rede.
# ============================================================================
def _fake_opener(dados, falhar_apos=None):
    """Cria um 'abrir' que serve `dados` a partir de `inicio` (Range). Se `falhar_apos` é dado,
    a PRIMEIRA chamada corta a leitura após N bytes (simula queda); as próximas servem normal."""
    import io
    estado = {"chamadas": 0}
    class _F(io.BytesIO):
        def __init__(self, buf, corta):
            super().__init__(buf); self._corta = corta; self._lido = 0
        def read(self, n=-1):
            if self._corta is not None and self._lido >= self._corta:
                raise OSError("conexão caiu (simulado)")
            b = super().read(n)
            self._lido += len(b)
            if self._corta is not None and self._lido > self._corta:
                # devolve só até o corte e marca para estourar na próxima
                excesso = self._lido - self._corta
                self._lido = self._corta
                return b[:len(b) - excesso]
            return b
    def abrir(url, inicio=0, timeout=120):
        estado["chamadas"] += 1
        corta = falhar_apos if (estado["chamadas"] == 1 and falhar_apos is not None) else None
        buf = dados[inicio:]
        return _F(buf, corta), len(buf), (inicio > 0)
    return abrir, estado

def test_downloader_sucesso_simples(tmp_path):
    from resources import downloader as dl
    dados = b"OPENROTAS" * 1000
    abrir, _ = _fake_opener(dados)
    out = tmp_path / "f.bin"
    eventos = []
    r = dl.baixar("http://x/f", out, progresso=lambda e: eventos.append(e), abrir=abrir)
    assert r["ok"] and r["bytes"] == len(dados) and out.read_bytes() == dados
    assert any(e.get("fase") == "baixando" for e in eventos)
    assert any(e.get("fase") == "concluido" for e in eventos)

def test_downloader_retoma_apos_queda(tmp_path):
    from resources import downloader as dl
    dados = b"x" * 1000
    abrir, estado = _fake_opener(dados, falhar_apos=400)   # cai após 400 bytes na 1ª tentativa
    out = tmp_path / "f.bin"
    r = dl.baixar("http://x/f", out, abrir=abrir, dormir=lambda s: None)
    assert r["ok"] and out.read_bytes() == dados
    assert r["tentativas"] == 2 and estado["chamadas"] == 2   # 2ª chamada retomou via Range

def test_downloader_sha256_rejeita(tmp_path):
    from resources import downloader as dl
    dados = b"conteudo"
    abrir, _ = _fake_opener(dados)
    out = tmp_path / "f.bin"
    r = dl.baixar("http://x/f", out, sha256="0" * 64, abrir=abrir, dormir=lambda s: None)
    assert r["ok"] is False and not out.exists()           # hash errado → descartado, sem retry infinito
    assert r["tentativas"] == 1

def test_downloader_eof_precoce_nao_finaliza_truncado(tmp_path):
    # Regressão do "verifique conexão": um servidor/proxy que corta o stream SEM erro (read()
    # devolve vazio antes do Content-Length) NÃO pode virar "download ok" — senão a parte truncada
    # corromperia o grafo concatenado. Deve tratar como falha e RETOMAR (Range) até completar.
    import io
    from resources import downloader as dl
    dados = b"Z" * 1000
    estado = {"n": 0}
    def abrir(url, inicio=0, timeout=120):
        estado["n"] += 1
        buf = dados[inicio:]
        if estado["n"] == 1:
            # declara o total cheio, mas entrega só 400 bytes e dá EOF (sem levantar)
            return io.BytesIO(buf[:400]), len(buf), (inicio > 0)
        return io.BytesIO(buf), len(buf), (inicio > 0)
    out = tmp_path / "f.bin"
    r = dl.baixar("http://x/f", out, abrir=abrir, dormir=lambda s: None, tentativas=4)
    assert r["ok"] is True
    assert out.read_bytes() == dados          # completou via retomada, sem truncar
    assert estado["n"] >= 2                    # a 1ª (truncada) não foi aceita; retomou

def test_downloader_sha256_aceita(tmp_path):
    import hashlib
    from resources import downloader as dl
    dados = b"conteudo-ok-123"
    abrir, _ = _fake_opener(dados)
    out = tmp_path / "f.bin"
    r = dl.baixar("http://x/f", out, sha256=hashlib.sha256(dados).hexdigest(), abrir=abrir)
    assert r["ok"] and out.read_bytes() == dados

def test_downloader_falha_persistente_desiste(tmp_path):
    from resources import downloader as dl
    def abrir(url, inicio=0, timeout=120):
        raise OSError("sem rede")
    r = dl.baixar("http://x/f", tmp_path / "f.bin", abrir=abrir, tentativas=3, dormir=lambda s: None)
    assert r["ok"] is False and r["tentativas"] == 3 and "falha" in r["detalhe"].lower()

def test_downloader_formatadores():
    from resources import downloader as dl
    assert dl.humano_velocidade(1024 * 1024).endswith("MB/s")
    assert dl.humano_eta(90) == "1min 30s" and dl.humano_eta(-1) == "—"
    assert dl.humano_eta(3700).endswith("min")

def test_prov_bases_ausentes_e_reassembly():
    from resources import provisionamento as prov
    # neste repo as bases nacionais já estão montadas → nenhuma ausente (exceto grafo, excluído)
    assert isinstance(prov.bases_ausentes(), list)
    r = prov._reassemblar_bigparts()
    assert isinstance(r, dict) and "ok" in r

def test_osrm_baixar_partes_ainda_funciona_via_downloader(tmp_path):
    # regressão: a nova trilha (downloader por parte) ainda concatena partes file:// corretamente
    import io, tarfile, gzip
    from engines import osrm_manager as osrm
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w") as t:
        dado = b"G" * 3000
        info = tarfile.TarInfo("brazil-latest.osrm"); info.size = len(dado)
        t.addfile(info, io.BytesIO(dado))
    targz = gzip.compress(buf.getvalue())
    srv = tmp_path / "srv"; srv.mkdir()
    meio = len(targz) // 2
    (srv / "g.tar.gz.part00").write_bytes(targz[:meio])
    (srv / "g.tar.gz.part01").write_bytes(targz[meio:])
    base = (srv / "g.tar.gz.part").as_uri()
    out = tmp_path / "grafo.tar.gz"
    assert osrm._baixar_partes(base, out) is True
    import gzip as _gz
    assert out.exists() and out.read_bytes() == targz


# ------------------------------- janela (GUI nativa) -------------------------------
# Testes HEADLESS-SAFE: nunca instanciam QApplication (não há display aqui). Cobrem as
# funções puras e o contrato de disponibilidade/fallback.
def test_janela_disponivel_e_booleano():
    assert isinstance(janela.disponivel(), bool)

def test_janela_pasta_downloads_resolve_e_existe(monkeypatch, tmp_path):
    # O destino dos downloads do app deve resolver para uma pasta real (nunca vazio/levanta).
    p = janela._pasta_downloads()
    assert isinstance(p, str) and p and os.path.isdir(p)
    # Chamar de novo continua resolvendo para uma pasta real (idempotente, nunca levanta).
    p2 = janela._pasta_downloads()
    assert isinstance(p2, str) and os.path.isdir(p2)

def test_janela_css_premium_tem_tokens_resolvidos():
    css = janela._css_premium()
    assert "@BRAND@" not in css and "@BG0@" not in css        # tokens substituídos
    assert janela.MARCA["brand"] in css                       # cor da marca presente
    assert "::-webkit-scrollbar" in css                       # scrollbar premium
    assert "stToolbar" in css                                 # esconde a barra do Streamlit

def test_janela_splash_html_bem_formado():
    html = janela._html_splash()
    assert html.lstrip().lower().startswith("<!doctype html>")
    assert "@TEXTO@" not in html and "@BRAND2@" not in html   # tokens substituídos
    assert "OpenRotas" in html and "Motor Nacional" in html

def test_janela_js_string_escapa_com_seguranca():
    import json
    bruto = 'quebra"aspas</style>\n<script>'
    assert json.loads(janela._js_string(bruto)) == bruto      # round-trip seguro


# ------------------------- config padrão embutida (turnkey) -------------------------
def test_config_ofuscacao_roundtrip():
    import gerar_config_padrao as gen
    original = '{"SUPABASE_URL":"https://x.supabase.co","SUPABASE_ANON_KEY":"sb_publishable_abc"}'
    blob = gen._ofuscar(original)
    assert blob != original.encode("utf-8")            # não é texto puro
    assert cfg._desofuscar(blob) == original           # mesmo salt → volta idêntico
    assert gen._OFUSCA_SALT == cfg._OFUSCA_SALT         # gerador e runtime casam

def test_config_eh_placeholder():
    assert cfg._eh_placeholder("https://SEUPROJETO.supabase.co")
    assert cfg._eh_placeholder("eyJ...SUA_CHAVE")
    assert cfg._eh_placeholder("") and cfg._eh_placeholder("   ")
    assert not cfg._eh_placeholder("https://meuprojetoreal0001.supabase.co")
    assert not cfg._eh_placeholder("sb_publishable_realkey")

def test_merge_placeholder_cai_no_embutido():
    d = {"SUPABASE_URL": "https://real.supabase.co", "SUPABASE_ANON_KEY": "sb_publishable_real"}
    u = {"SUPABASE_URL": "https://SEUPROJETO.supabase.co", "SUPABASE_ANON_KEY": "eyJ...SUA_CHAVE", "offline": False}
    eff = cfg._mesclar_config(d, u)
    assert eff["SUPABASE_URL"] == "https://real.supabase.co"      # placeholder → embutido
    assert eff["SUPABASE_ANON_KEY"] == "sb_publishable_real"
    assert eff["offline"] is False                                 # outros campos do usuário entram

def test_merge_usuario_real_vence_sobre_embutido():
    d = {"SUPABASE_URL": "https://real.supabase.co", "SUPABASE_ANON_KEY": "sb_publishable_real"}
    u = {"SUPABASE_URL": "https://meuproprio.supabase.co"}
    eff = cfg._mesclar_config(d, u)
    assert eff["SUPABASE_URL"] == "https://meuproprio.supabase.co"  # usuário real vence
    assert eff["SUPABASE_ANON_KEY"] == "sb_publishable_real"        # ausente → embutido

def test_gerador_recusa_service_role(monkeypatch, tmp_path):
    import gerar_config_padrao as gen
    monkeypatch.setenv("ORK_SUPABASE_URL", "https://x.supabase.co")
    monkeypatch.setenv("ORK_SUPABASE_ANON_KEY", "sb_secret_PERIGO")
    assert gen.main() == 2        # trava: nunca embutir secret/service_role

def test_gerador_mapeia_motor_dedicado():
    # Option A: o motor de rotas DEDICADO (acelera o lote) é embutível via secrets do build.
    import gerar_config_padrao as gen
    assert gen._MAPA.get("ORK_OSRM_URL") == "OSRM_URL"
    assert gen._MAPA.get("ORK_VALHALLA_URL") == "VALHALLA_URL"
    assert gen._MAPA.get("ORK_GRAPHHOPPER_URL") == "GRAPHHOPPER_URL"
    # a URL de um OSRM próprio não é um segredo (service_role) — não pode cair na trava
    assert not any(p in "http://osrm.meuservidor.com:5000".lower() for p in gen._PROIBIDO)

def test_osrm_url_embutido_sobrevive_e_sinaliza_motor_proprio():
    # Um OSRM_URL embutido (default turnkey) chega ao app quando o usuário não preencheu o seu,
    # e aponta para uma instância PRÓPRIA (não o público) — o que deixa o lote rápido.
    d = {"SUPABASE_URL": "https://real.supabase.co", "SUPABASE_ANON_KEY": "sb_publishable_real",
         "OSRM_URL": "http://osrm-brasil.meuservidor.com:5000"}
    eff = cfg._mesclar_config(d, {"offline": False})
    assert eff["OSRM_URL"] == "http://osrm-brasil.meuservidor.com:5000"   # embutido entra
    assert "project-osrm.org" not in eff["OSRM_URL"].lower()              # instância própria

def test_osrm_url_do_usuario_vence_sobre_embutido():
    d = {"OSRM_URL": "http://padrao.turnkey:5000"}
    eff = cfg._mesclar_config(d, {"OSRM_URL": "http://meu-osrm-local:5000"})
    assert eff["OSRM_URL"] == "http://meu-osrm-local:5000"   # usuário real vence o embutido


# ------------------- motor NATIVO embarcado (sem Docker) — turnkey local -------------------
def test_osrm_bin_locator_e_disponibilidade(tmp_path, monkeypatch):
    # sem binário → indisponível; com OSRM_ROUTED_BIN apontando p/ um arquivo → localiza
    monkeypatch.delenv("OSRM_ROUTED_BIN", raising=False)
    # garante que não ache um binário dev acidental
    monkeypatch.setattr(osrm, "_nome_binario", lambda: "osrm-routed-inexistente-xyz")
    assert osrm._bin_osrm_routed() is None
    assert osrm.motor_nativo_disponivel() is False
    fake = tmp_path / "osrm-routed"
    fake.write_text("#!/bin/sh\n")
    monkeypatch.setenv("OSRM_ROUTED_BIN", str(fake))
    assert osrm._bin_osrm_routed() == str(fake)
    assert osrm.motor_nativo_disponivel() is True

def test_osrm_resolver_local_sem_grafo_cai_no_publico(tmp_path, monkeypatch):
    # modo 'local' sem grafo → inativo (não serve), app cai no público (zero regressão)
    fake = tmp_path / "osrm-routed"; fake.write_text("x")
    monkeypatch.setenv("OSRM_ROUTED_BIN", str(fake))
    r = osrm.resolver({"mode": "local", "graph_path": str(tmp_path / "nao_existe.osrm")})
    assert r.modo == "local" and r.ativo is False and r.url is None

def test_osrm_num_defensivo():
    assert osrm._num(None, 5000) == 5000              # JSON null → default
    assert osrm._num("", 5000) == 5000                # string vazia → default
    assert osrm._num("8080", 5000) == 8080            # string numérica ok
    assert osrm._num(90, 30, float) == 90.0
    assert osrm._num("x", 90, float) == 90.0          # lixo → default

def test_osrm_resolver_porta_invalida_nao_levanta(tmp_path, monkeypatch):
    # osrm.port null/"" no desktop.json NÃO pode fazer resolver() levantar (perderia o fallback
    # para o OSRM público). Deve devolver inativo, sem exceção.
    monkeypatch.setattr(osrm, "_bin_osrm_routed", lambda: None)
    for porta in (None, "", "abc"):
        r = osrm.resolver({"mode": "local", "graph_path": str(tmp_path / "g.osrm"), "port": porta})
        assert r.ativo is False and r.url is None      # sem levantar

def test_osrm_resolver_auto_sem_motor_cai_no_publico(monkeypatch):
    # modo 'auto' sem binário nativo e sem Docker → inativo, público (nunca levanta)
    monkeypatch.setattr(osrm, "_bin_osrm_routed", lambda: None)
    monkeypatch.setattr(osrm, "_docker_disponivel", lambda: False)
    r = osrm.resolver({"mode": "auto"})
    assert r.modo == "auto" and r.ativo is False and r.url is None

def test_osrm_resolver_local_sem_binario_cai_no_publico(tmp_path, monkeypatch):
    # modo 'local' COM grafo mas SEM binário embarcado → inativo (fallback seguro)
    g = tmp_path / "brazil.osrm"; g.write_text("grafo")
    monkeypatch.setattr(osrm, "_bin_osrm_routed", lambda: None)
    r = osrm.resolver({"mode": "local", "graph_path": str(g)})
    assert r.modo == "local" and r.ativo is False and r.url is None

def test_osrm_encerrar_motor_nativo_encerra_processo_sem_docker():
    # encerrar() de um motor NATIVO (sem container) termina o processo e NÃO chama docker
    class _FakeProc:
        def __init__(self): self.terminado = False
        def terminate(self): self.terminado = True
        def wait(self, timeout=None): return 0
        def kill(self): pass
    p = _FakeProc()
    res = osrm.ResultadoMotor(url="http://127.0.0.1:5000", modo="local", ativo=True,
                              gerenciado=True, processo=p, nome="")   # nome vazio = nativo
    osrm.encerrar(res)
    assert p.terminado is True

def test_osrm_encerrar_nao_gerenciado_e_noop():
    res = osrm.ResultadoMotor(url=None, modo="off", ativo=False, gerenciado=False)
    osrm.encerrar(res)   # não levanta


# ------------------------- camada premium ESTÁTICA (desktop) -------------------------
def test_janela_css_premium_e_estatica_sem_flicker():
    # Regressão do "piscar o tempo inteiro": a camada premium NÃO pode ter animação contínua
    # (aurora infinita), NEM animação de entrada que o Streamlit reexecuta a cada rerun
    # (ork-rise), NEM blur de tela cheia animado — tudo causava repaint constante no QtWebEngine.
    css = janela._css_premium()
    for proibido in ("ork-aurora", "ork-rise", "@keyframes", "animation:", "preserve-3d"):
        assert proibido not in css, "efeito que pisca voltou ao CSS premium: " + proibido
    # mas a PROFUNDIDADE estática e o acabamento continuam presentes
    assert "box-shadow" in css
    assert "::-webkit-scrollbar" in css
    assert "stToolbar" in css
    assert "@media (prefers-reduced-motion" in css   # ainda respeita acessibilidade

def test_janela_js_premium_sem_tilt_3d():
    # O tilt 3D seguia o mouse reescrevendo transform (repaint contínuo + efeito 3D indesejado):
    # foi removido. O JS fica inerte e idempotente.
    js = janela._js_premium()
    assert "__orkPremium" in js            # guarda de idempotência (inerte)
    assert "mousemove" not in js           # sem listener de movimento
    assert "perspective(" not in js        # sem inclinação 3D
    assert "try" in js and "catch" in js   # defensivo


# ------------------------- notificação nativa (desktop) -------------------------
def test_notificacao_importa_e_e_noop_fora_do_windows():
    from resources import notificacao
    # fora do Windows (ambiente de teste) é no-op seguro, nunca levanta
    assert isinstance(notificacao.disponivel(), bool)
    assert notificacao.notificar("t", "m") in (True, False)

def test_streamlit_hook_notificacao_inerte_sem_env(monkeypatch):
    # o hook do app só dispara com OPENROTAS_DESKTOP=1; sem a env, retorna cedo (inerte na web)
    import importlib.util, os
    monkeypatch.delenv("OPENROTAS_DESKTOP", raising=False)
    _p = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "streamlit_app.py")
    src = open(_p, encoding="utf-8").read()
    assert "_notificar_conclusao_desktop" in src and 'OPENROTAS_DESKTOP") != "1"' in src
